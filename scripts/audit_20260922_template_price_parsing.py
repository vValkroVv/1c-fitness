#!/usr/bin/env python3
"""Read source XLSX values/formats and compare saved Fitbase catalog prices."""
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output/20260922_fitbase_live_audit/price_parsing'
AUDIT = OUT.parent
SOURCE = ROOT / 'output/20260921_fitbase_for_customer'
CLUB_LABELS = {
    'Stolitsa_allstatuses':'Столица', 'Stolitsa_archive':'Столица',
    'Gog_active':'Гоголевский', 'Gog_archive':'Гоголевский',
    'Rovio_active_verified':'Ровио', 'Rovio_archive':'Ровио',
    'Industrial_active':'Промышленная', 'Industrial_archive':'Промышленная',
}

def dump(name, value):
    (OUT/name).write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n')

def source_cells(path):
    book = load_workbook(path,read_only=True,data_only=False)
    sheet = book.active
    rows = iter(sheet.iter_rows())
    headers = [c.value for c in next(rows)]
    next(rows)
    price_idx, name_idx = headers.index('price'), headers.index('name')
    result = []
    for row in rows:
        cell = row[price_idx]
        result.append({'name':row[name_idx].value,'cell':cell.coordinate,
                       'row':cell.row,'price':cell.value,'data_type':cell.data_type,
                       'python_type':type(cell.value).__name__,'number_format':cell.number_format})
    book.close()
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        ns = {'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        xml = {c.get('r'):c for c in root.findall('.//x:c',ns)}
        for r in result:
            cell = xml[r['cell']]
            r['xml_cell_type'] = cell.get('t')
            value = cell.find('x:v',ns)
            r['xml_value'] = value.text if value is not None else None
            r['literal_whitespace'] = bool(r['xml_value'] and re.search(r'\s',r['xml_value']))
    return {'file':path.name,'sheet':sheet.title,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
            'rows':result,'formats':dict(Counter(r['number_format'] for r in result)),
            'cell_types':dict(Counter(r['data_type'] for r in result))}

def main():
    OUT.mkdir(exist_ok=True)
    original = json.loads((AUDIT/'source_summary.json').read_text())
    saved = json.loads((AUDIT/'catalog_reconciliation_final.json').read_text())
    evidence = json.loads((AUDIT/'browser_evidence_part2.json').read_text())
    clubs = defaultdict(set)
    for selection in evidence['catalogs']:
        if selection['label'] not in CLUB_LABELS:
            continue
        for row in selection['rows']:
            for link in row['links']:
                if '/ticket/update?' in link['href']:
                    clubs[link['href']].add(CLUB_LABELS[selection['label']])
    all_results, sources, summaries = {}, {}, {}
    for kind, table in [('membership','membership_templates'),('service','service_templates')]:
        source = source_cells(SOURCE/original[table]['file'])
        assert source['sha256'] == original[table]['sha256']
        assert all(r['data_type']=='n' and not r['literal_whitespace'] for r in source['rows'])
        sources[kind] = source
        names = {r['name']:r for r in source['rows']}
        result = []
        for path, row in saved[table].items():
            cells = row['cells']
            name = cells[2] if kind=='membership' else cells[-8]
            raw = cells[5] if kind=='membership' else cells[-5]
            prices = [float(v) for v in raw.split() if re.fullmatch(r'\d+(\.\d+)?',v)]
            expected, basis = names.get(name), 'exact_name'
            candidates = []
            if expected is None:
                candidates = [r for n,r in names.items() if len(name)>=50 and n.startswith(name)]
                # A missing word could identify a different offer; do not equate it.
                if len(candidates)==1 and len(candidates[0]['name'])-len(name)<=4:
                    expected, basis = candidates[0], 'unique_short_truncated_suffix'
            item = {'url':'https://fitnes-imperiya.fitbase.io'+path,'live_name':name,
                    'clubs':sorted(clubs[path]),'live_prices':prices,'raw_catalog_price':raw,
                    'match_basis':basis if expected else None}
            count_cell=next(c for c in cells if 'Действующие:' in c and 'Всего:' in c)
            active=int(re.search(r'Действующие:\s*(\d+)',count_cell)[1])
            total=int(re.search(r'Всего:\s*(\d+)',count_cell)[1])
            assert total>=active
            item.update({'catalog_active_purchases':active,'catalog_total_purchases':total})
            if expected:
                value = expected['price']
                status = ('equal' if value in prices else
                          'thousands_prefix' if value>=1000 and value//1000 in prices else
                          'zero' if value>0 and prices and set(prices)=={0} else
                          'one' if prices==[1] else 'other')
                item.update({'source_name':expected['name'],'source_cell':expected['cell'],
                             'source_price':value,'source_number_format':expected['number_format'],
                             'source_xml_value':expected['xml_value'],'classification':status})
            else:
                item['classification']='unmatched_name'
                if candidates:
                    item['possible_source_names']=[r['name'] for r in candidates]
            result.append(item)
        all_results[kind]=result
        summaries[kind]={'source_rows':len(source['rows']),'live_template_ids':len(result),
                         'classification':dict(Counter(r['classification'] for r in result)),
                         'matched_source_names':len({r['source_name'] for r in result if r.get('source_name')}),
                         'by_match_basis':{basis:dict(Counter(r['classification'] for r in result
                                                            if r['match_basis']==basis))
                                           for basis in ['exact_name','unique_short_truncated_suffix',None]},
                         'multiple_catalog_prices':sum(len(r['live_prices'])>1 for r in result),
                         'severe_below_100':sum(r.get('source_price',0)>=1000 and bool(r['live_prices'])
                                              and max(r['live_prices'])<100 for r in result)}
        affected=[r for r in result if r['classification'] not in ['equal','unmatched_name']]
        summaries[kind]['affected_catalog_counts']={
            'template_ids':len(affected),
            'with_active_purchases':sum(r['catalog_active_purchases']>0 for r in affected),
            'without_active_purchases':sum(r['catalog_active_purchases']==0 for r in affected),
            'active_purchases':sum(r['catalog_active_purchases'] for r in affected),
            'other_purchases':sum(r['catalog_total_purchases']-r['catalog_active_purchases'] for r in affected),
        }
        if kind=='membership':
            summaries[kind]['by_club']={cl:dict(Counter(r['classification'] for r in result if cl in r['clubs']))
                                        for cl in sorted(set(CLUB_LABELS.values()))}
        fields=['url','clubs','live_name','source_name','source_cell','source_price',
                'source_number_format','source_xml_value','live_prices','classification','match_basis',
                'catalog_active_purchases','catalog_total_purchases']
        with (OUT/f'{kind}_all_prices.csv').open('w',encoding='utf-8-sig',newline='') as f:
            writer=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');writer.writeheader()
            for r in result:
                writer.writerow({**r,'clubs':', '.join(r['clubs']),
                                 'live_prices':', '.join(str(v) for v in r['live_prices'])})
    groups=defaultdict(lambda:defaultdict(list))
    for r in all_results['membership']:
        for cl in r['clubs']:
            groups[r['live_name']][cl].append(r)
    cross=[]
    for name,g in groups.items():
        if set(g)!=set(CLUB_LABELS.values()) or any(len(v)!=1 or len(v[0]['live_prices'])!=1 for v in g.values()):
            continue
        values={cl:v[0]['live_prices'][0] for cl,v in g.items()}
        if (values['Гоголевский']==values['Ровио'] and values['Промышленная']==values['Столица']
            and values['Гоголевский']>=1000 and values['Промышленная']==values['Гоголевский']//1000):
            cross.append({'name':name,'prices':values,'urls':{cl:v[0]['url'] for cl,v in g.items()}})
    summaries['cross_club_truncation']={'names':len(cross),'low_price_template_ids':2*len(cross),
        'low_clubs':['Промышленная','Столица'],'full_price_clubs':['Гоголевский','Ровио']}
    cross_urls={r['urls'][cl] for r in cross for cl in ['Промышленная','Столица']}
    low_rows=[r for r in all_results['membership'] if r['url'] in cross_urls]
    summaries['cross_club_truncation'].update({
        'with_active_purchases':sum(r['catalog_active_purchases']>0 for r in low_rows),
        'without_active_purchases':sum(r['catalog_active_purchases']==0 for r in low_rows),
        'active_purchases':sum(r['catalog_active_purchases'] for r in low_rows),
        'other_purchases':sum(r['catalog_total_purchases']-r['catalog_active_purchases'] for r in low_rows),
    })
    assert sum(summaries['membership']['classification'].values())==436
    assert sum(summaries['service']['classification'].values())==177
    dump('source_cells.json',sources)
    dump('all_comparisons.json',all_results)
    dump('cross_club_truncation.json',cross)
    dump('summary.json',summaries)
    print(json.dumps(summaries,ensure_ascii=False,indent=2))

if __name__=='__main__':
    main()
