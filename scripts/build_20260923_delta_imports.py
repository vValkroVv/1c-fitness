#!/usr/bin/env python3
"""Rebuild FitBase import deltas from the true September 20/23 backup cuts."""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import re
import subprocess
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "end-to-end-xlsx"
OLD = PACKAGE / "work/20260922_final"
NEW = PACKAGE / "work/20260923_restore_verified"
OUT = ROOT / "output/20260923_delta_from_20260920"
REPORTS = OUT / "reports/imports"
SUPPLEMENT = ROOT / "output/20260921_phone_dedup_supplement_371"
SUPPLEMENT_WORK = PACKAGE / "work/20260921_phone_dedup_supplement_371"
OLD_CUTOFF = "2026-09-20 20:12:12"
NEW_CUTOFF = "2026-09-23 23:36:39"
SUPPLEMENT_CUTOFF = "2026-09-21 20:12:12"


def module(path: str, name: str):
    spec = importlib.util.spec_from_file_location(name, PACKAGE / "scripts" / path)
    obj = importlib.util.module_from_spec(spec)
    import sys
    sys.modules[name] = obj
    spec.loader.exec_module(obj)
    return obj


def csv_rows(path: Path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv(path: Path, rows: list[dict], fields: list[str]):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def workbook_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    it = wb.active.iter_rows(values_only=True)
    headers = [str(x or "") for x in next(it)]
    next(it)
    rows = [dict(zip(headers, values)) for values in it if any(x is not None for x in values)]
    wb.close()
    return rows


def normalized_phones(raw: str):
    result = set()
    for part in re.split(r"[,;]", raw or ""):
        digits = re.sub(r"\D", "", part)
        if len(digits) == 10:
            result.add("7" + digits)
        elif len(digits) == 11 and digits[0] in "78":
            result.add("7" + digits[1:])
    return result


def val(value):
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()[:10]
    return str(value).strip()


def service_signature(row):
    return tuple(val(row.get(k)) for k in (
        "service_id", "client_id", "phone", "client_fio", "service_name",
        "create_date", "payment_date", "activation_date", "end_date", "count",
        "visits_left", "price", "amount_of_payment", "payment_left",
        "type_of_payment", "manager", "филиал"))


def in_window(value):
    if not value or len(value) < 19:
        return False
    datetime.strptime(value[:19], "%Y-%m-%d %H:%M:%S")
    return OLD_CUTOFF < value[:19] <= NEW_CUTOFF


def check_supplement_baseline():
    """Read the prepared supplement as a separate prior delivery, at its own cutoff."""
    meta = json.loads((SUPPLEMENT_WORK / "summary.json").read_text())
    if meta["backup_finish_at"] != OLD_CUTOFF or meta["cutoff_at"] != SUPPLEMENT_CUTOFF:
        raise RuntimeError(f"Supplement backup/cutoff mismatch: {meta}")
    clients = csv_rows(SUPPLEMENT_WORK / "selected_clients.csv")
    leads = workbook_rows(SUPPLEMENT / "fitbase_active_clients_import_zayavki_20260921_all_funnels.xlsx")
    members = workbook_rows(SUPPLEMENT / "fitbase_import_abonementy_clientov_20260921.xlsx")
    services = workbook_rows(SUPPLEMENT / "fitbase_import_uslugi_clientov_20260921.xlsx")
    if (len(clients), len(leads), len(members), len(services)) != (371, 371, 879, 10):
        raise RuntimeError("Supplement row counts changed")
    if {r["client_id"] for r in clients} != {val(r["client_id"]) for r in leads}:
        raise RuntimeError("Supplement selected clients differ from published XLSX")
    for rows, key in ((leads, "client_id"), (members, "contract_id"), (services, "service_id")):
        if len({val(r[key]) for r in rows}) != len(rows):
            raise RuntimeError(f"Duplicate supplement {key}")
    return clients, leads, members, services


def verify_existing_refuser_buyers(client_ids, old_members, current_stage_by_id):
    """Tie prior membership-only customers to one active Fitbase card each."""
    old_by_id = defaultdict(list)
    for row in old_members:
        client_id = val(row.get("client_id"))
        if client_id in client_ids and val(row.get("tag")) == "отказники":
            old_by_id[client_id].append(row)
    wanted_phones = set()
    for rows in old_by_id.values():
        for row in rows:
            wanted_phones.update(normalized_phones(val(row.get("phone"))))
    api_by_phone = defaultdict(dict)
    snapshot = ROOT / "output/20260923_fitbase_photo_audit/api_snapshot.jsonl"
    metadata = json.loads((snapshot.parent / "api_snapshot_metadata.json").read_text())
    if metadata.get("completeness_check") != "passed" or metadata.get("records_written") != metadata.get("unique_ids"):
        raise RuntimeError("Fitbase API snapshot completeness is unverified")
    with snapshot.open(encoding="utf-8") as handle:
        for line in handle:
            card = json.loads(line)
            for contact in card.get("contacts") or []:
                for phone in normalized_phones(val(contact.get("contact"))) & wanted_phones:
                    api_by_phone[phone][card["id"]] = card
    audit = []
    confirmed = set()
    for client_id in sorted(client_ids):
        old_rows = old_by_id.get(client_id, [])
        if not old_rows:
            continue
        stage = current_stage_by_id.get(client_id)
        old_phones = set().union(*(normalized_phones(val(r.get("phone"))) for r in old_rows))
        current_phones = normalized_phones(stage.get("phones", "")) if stage else set()
        matching_phones = old_phones & current_phones
        matches = {card_id: card for phone in matching_phones for card_id, card in api_by_phone[phone].items()}
        active = {card_id: card for card_id, card in matches.items() if not card.get("is_archive")}
        old_cards = {val(r.get("card")) for r in old_rows if val(r.get("card"))}
        card = next(iter(active.values())) if len(active) == 1 else None
        tags = set(val(v) for v in (card.get("tags") or {}).values()) if card else set()
        card_number = val(card.get("card")) if card else ""
        if not stage or not matching_phones:
            status = "prior_and_current_phone_mismatch"
        elif len(matches) != 1 or len(active) != 1:
            status = "fitbase_card_not_unique_or_inactive"
        elif "отказники" not in tags:
            status = "fitbase_refuser_tag_missing"
        elif old_cards and card_number not in old_cards:
            status = "fitbase_card_number_mismatch"
        else:
            status = "existing_target_confirmed"
            confirmed.add(client_id)
        audit.append({
            "client_id": client_id,
            "client_fio": val(old_rows[0].get("client_fio")),
            "normalized_phone": ", ".join(sorted(matching_phones)),
            "prior_membership_rows": len(old_rows),
            "prior_tag": "отказники",
            "prior_card": ", ".join(sorted(old_cards)),
            "fitbase_match_count": len(matches),
            "fitbase_active_match_count": len(active),
            "fitbase_id": card["id"] if card else "",
            "fitbase_name": " ".join(val(card.get(key)) for key in ("surname", "name", "patronymic")).strip() if card else "",
            "fitbase_card": card_number,
            "status": status,
        })
    return confirmed, audit, metadata["completed_at_utc"]


def workbook_matches(path: Path, headers: list[str], rows: list[dict]):
    if not path.exists():
        return False
    actual = workbook_rows(path)
    return len(actual) == len(rows) and all(
        list(found) == list(headers) and all(val(found.get(h)) == val(wanted.get(h)) for h in headers)
        for found, wanted in zip(actual, rows, strict=True)
    )


def build():
    if not (NEW / "status.json").exists():
        raise RuntimeError("New full pipeline staging is not ready")
    status = json.loads((NEW / "status.json").read_text())
    contract = status["cutoff_contract"]
    if contract["backup_finish_at"] != NEW_CUTOFF or contract["cutoff_at"] != NEW_CUTOFF:
        raise RuntimeError(f"New backup cutoff mismatch: {contract}")
    old_meta = csv_rows(OLD / "owner/staging/staging_run_metadata.csv")
    old_values = {val(v) for row in old_meta for v in row.values()}
    if not any(v.startswith("2026-09-22 20:12:12") for v in old_values):
        raise RuntimeError("Old synthetic delivery metadata changed; check baseline")
    if not any(v.startswith(OLD_CUTOFF) for v in old_values):
        raise RuntimeError("Old true BackupFinishDate mismatch")
    reports = REPORTS
    reports.mkdir(parents=True, exist_ok=True)
    supplement_clients, supplement_leads, supplement_members, supplement_services = check_supplement_baseline()
    source = NEW / "owner/staging/final_funnel_clients.csv"
    old_source = OLD / "owner/staging/final_funnel_clients.csv"
    old_clients = csv_rows(old_source)
    new_clients = csv_rows(source)
    for row in new_clients:
        if row.get("cutoff_date") != "2026-09-23":
            raise RuntimeError("Client layer cutoff mismatch")
    old_refs = {r["client_ref"] for r in old_clients} | {r["client_ref"] for r in supplement_clients}
    raw_new_leads = [r for r in new_clients if r["client_ref"] not in old_refs]
    current_published = workbook_rows(NEW / "owner/fitbase_active_clients_import_zayavki_20260923_all_funnels.xlsx")
    old_published = workbook_rows(ROOT / "output/20260922_fitbase_for_customer/fitbase_active_clients_import_zayavki_20260922_all_funnels.xlsx")
    current_by_id = {val(r["client_id"]): r for r in current_published}
    old_published_ids = {val(r["client_id"]) for r in old_published + supplement_leads}
    old_by_id = defaultdict(list)
    old_by_phone = defaultdict(list)
    for row in {r["client_ref"]: r for r in old_clients + supplement_clients}.values():
        old_by_id[row["client_id"]].append(row)
        for phone in normalized_phones(row.get("phones", "")):
            old_by_phone[phone].append(row)
    collisions = []
    for row in raw_new_leads:
        for prior in old_by_id.get(row["client_id"], []):
            collisions.append({"new_client_ref": row["client_ref"], "new_client_id": row["client_id"], "new_phone": row["phones"], "collision_kind": "client_id", "old_client_ref": prior["client_ref"], "old_client_id": prior["client_id"], "old_phone": prior["phones"]})
        for phone in normalized_phones(row.get("phones", "")):
            for prior in old_by_phone.get(phone, []):
                collisions.append({"new_client_ref": row["client_ref"], "new_client_id": row["client_id"], "new_phone": row["phones"], "collision_kind": "normalized_phone:" + phone, "old_client_ref": prior["client_ref"], "old_client_id": prior["client_id"], "old_phone": prior["phones"]})
    write_csv(reports / "new_lead_collisions.csv", collisions, ["new_client_ref", "new_client_id", "new_phone", "collision_kind", "old_client_ref", "old_client_id", "old_phone"])
    collision_refs = {r["new_client_ref"] for r in collisions}
    new_leads = []
    lead_scope = []
    seen_phones = {}
    new_phone_collisions = []
    for row in raw_new_leads:
        phones = normalized_phones(row.get("phones", ""))
        if not phones:
            disposition = "invalid_phone"
        elif row["client_ref"] in collision_refs or row["client_id"] in old_published_ids:
            disposition = "old_client_id_or_phone_collision"
        elif phones & seen_phones.keys():
            disposition = "new_client_phone_collision"
            for phone in sorted(phones & seen_phones.keys()):
                new_phone_collisions.append({"phone": phone, "excluded_client_ref": row["client_ref"], "excluded_client_id": row["client_id"], "kept_client_ref": seen_phones[phone]["client_ref"], "kept_client_id": seen_phones[phone]["client_id"]})
        else:
            disposition = "import"
            new_leads.append(row)
            seen_phones.update({phone: row for phone in phones})
        lead_scope.append({"client_ref": row["client_ref"], "client_id": row["client_id"], "funnel": row["funnel"], "in_current_full_export": "1" if row["client_id"] in current_by_id else "0", "disposition": disposition})
    write_csv(reports / "new_lead_duplicate_phone.csv", new_phone_collisions, ["phone", "excluded_client_ref", "excluded_client_id", "kept_client_ref", "kept_client_id"])
    funnel = module("17_build_part2_combined_xlsx.py", "delta_funnel")
    base = funnel.load_three_funnel_builder()
    base.assign_managers(new_leads, base.load_managers(PACKAGE / "config/managers_by_club.yml"))
    base.assign_branches(new_leads, base.load_branches(PACKAGE / "config/branches_by_club.yml"))
    new_leads = base.sort_rows(new_leads)
    labels = funnel.apply_fitbase_labels(new_leads, funnel.CUSTOMER_SINGLE_STAGE_MODE)
    lead_rows = []
    for raw, labeled in zip(new_leads, labels, strict=True):
        if raw["client_id"] in current_by_id:
            lead_rows.append({h: val(current_by_id[raw["client_id"]].get(h)) for h in base.MAIN_HEADERS})
        else:
            lead_rows.append({"client_id": raw["client_id"], "phone": raw["phones"], "client_fio": raw["client_fio"], "email": raw["email"], "funnel": labeled["funnel"], "funnel_step": labeled["funnel_step"], "budget": 0, "create_date": raw["create_date"], "manager": raw["manager"], "филиал": raw["branch"]})
    memberships = module("19_build_membership_import_xlsx.py", "delta_memberships")
    owner = NEW / "owner"
    owner_file = owner / "fitbase_active_clients_import_zayavki_20260923_all_funnels.xlsx"
    source_clients = memberships.read_source_clients(owner_file)
    source_clients.update(memberships.read_refuser_clients(owner / "csv/new_application_refusers.csv"))
    cards = memberships.read_cards(owner / "staging")
    member_facts = memberships.read_facts(NEW / "imports/staging/membership_import_facts.tsv")
    if any(r.get("cutoff_at") != NEW_CUTOFF for r in member_facts):
        raise RuntimeError("Membership layer cutoff mismatch")
    old_member_refs = {r["_subscription_ref"] for r in csv_rows(OLD / "imports/staging/membership_import_rows.csv") if r.get("_subscription_ref")}
    old_main_members = workbook_rows(ROOT / "output/20260922_fitbase_for_customer/fitbase_import_abonementy_clientov_20260922.xlsx")
    old_contract_ids = {val(r["contract_id"]) for r in old_main_members + supplement_members}
    selected_members = [r for r in member_facts if in_window(r.get("sale_datetime")) and r.get("subscription_ref") not in old_member_refs]
    if len({r["subscription_ref"] for r in selected_members}) != len(selected_members):
        raise RuntimeError("Duplicate new membership reference")
    excluded_refs, _ = memberships.find_contact_next_exclusions(selected_members)
    member_rows = []
    member_issues = []
    for fact in selected_members:
        if fact["subscription_ref"] in excluded_refs:
            continue
        mapped, _, issues, _, _ = memberships.build_rows(source_clients, cards, [fact], {})
        member_rows.extend(r for r in mapped if r.get("_subscription_ref"))
        member_issues.extend(issues)
    member_by_ref = {r["_subscription_ref"]: r for r in member_rows}
    member_exclusions = []
    for fact in selected_members:
        ref = fact["subscription_ref"]
        if ref not in member_by_ref:
            reason = "client_not_in_current_import" if fact.get("client_id") not in source_clients else "builder_business_exclusion"
            member_exclusions.append({"subscription_ref": ref, "sale_datetime": fact["sale_datetime"], "client_id": fact.get("client_id", ""), "reason": reason})
    services = module("23_build_services_import_xlsx.py", "delta_services")
    service_source_clients = services.read_source_clients(owner_file)
    service_facts = services.read_facts(NEW / "imports/staging/services_import_facts.tsv")
    extra_facts_path = reports / "extra_services_import_facts.tsv"
    if not extra_facts_path.exists():
        raise RuntimeError("Supplementary service fact export missing; run scripts/export_20260923_extra_services.py")
    extra_facts = services.read_facts(extra_facts_path)
    if any(r.get("cutoff_at") != NEW_CUTOFF for r in service_facts):
        raise RuntimeError("Service layer cutoff mismatch")
    if any(r.get("cutoff_at") != NEW_CUTOFF for r in extra_facts):
        raise RuntimeError("Supplementary service layer cutoff mismatch")
    configured_window = [r for r in service_facts if in_window(r.get("sale_datetime"))]
    extra_window = [r for r in extra_facts if in_window(r.get("sale_datetime"))]
    window_services = configured_window + extra_window
    extra_names = set(json.loads((reports / "extra_services_export_result.json").read_text())["extra_names"])
    if len(extra_names) != 9 or {r["service_name"] for r in extra_window} != extra_names:
        raise RuntimeError("The nine additional service names differ from the staged purchase facts")
    if extra_names & {r["service_name"] for r in configured_window}:
        raise RuntimeError("Additional and configured service catalogues overlap")
    extra_sale_keys = {(r["sale_doc_ref"], r["sale_line_no"]) for r in extra_window}
    keys = [(r["sale_doc_ref"], r["sale_line_no"]) for r in window_services]
    if len(set(keys)) != len(keys):
        raise RuntimeError("Duplicate service sale line composite key")
    old_service_rows = workbook_rows(ROOT / "output/20260922_fitbase_for_customer/fitbase_import_uslugi_clientov_20260922.xlsx") + supplement_services
    old_signatures = {service_signature(r) for r in old_service_rows}
    old_service_ids = {val(r["service_id"]) for r in old_service_rows}
    # The established mapper is reused, with each window sale eligible for selection.
    # Its normal active/five-history selection is inappropriate for a sale delta.
    eligible = [{**r, "is_active_on_cutoff": "1"} for r in window_services]
    service_names = sorted({r["service_name"] for r in eligible})
    manager_pools = services.manager_tools().load_managers(PACKAGE / "config/managers_by_club.yml")
    service_rows = []
    service_issues = []
    for fact in eligible:
        mapped, _, issues, _, _ = services.build_rows(service_source_clients, [fact["service_name"]], [fact], manager_pools=manager_pools)
        if len(mapped) != 1:
            raise RuntimeError(f"Service mapper lost sale {fact['sale_doc_ref']}:{fact['sale_line_no']}")
        service_rows.append(mapped[0])
        service_issues.extend(issues)
    known_target_ids = old_published_ids | {r["client_id"] for r in lead_rows}
    current_stage_by_id = {r["client_id"]: r for r in new_clients}
    buyer_ids = {r["client_id"] for r in member_rows} | {r["client_id"] for r in service_rows}
    old_refuser_ids = {val(r.get("client_id")) for r in old_main_members if val(r.get("tag")) == "отказники"}
    existing_candidates = (buyer_ids - known_target_ids) & old_refuser_ids
    existing_target_ids, existing_audit, snapshot_at = verify_existing_refuser_buyers(
        existing_candidates, old_main_members, current_stage_by_id
    )
    write_csv(reports / "existing_refuser_buyers_fitbase_audit.csv", existing_audit, [
        "client_id", "client_fio", "normalized_phone", "prior_membership_rows",
        "prior_tag", "prior_card", "fitbase_match_count", "fitbase_active_match_count",
        "fitbase_id", "fitbase_name", "fitbase_card", "status",
    ])
    known_target_ids.update(existing_target_ids)
    managers_by_club = base.load_managers(PACKAGE / "config/managers_by_club.yml")
    branches_by_club = base.load_branches(PACKAGE / "config/branches_by_club.yml")
    prerequisite_rows = []
    prerequisite_review = []
    for client_id in sorted(buyer_ids - known_target_ids):
        if client_id in existing_candidates:
            prerequisite_review.append({"client_id": client_id, "reason": "prior_refuser_fitbase_match_unconfirmed"})
            continue
        stage = current_stage_by_id.get(client_id)
        if not stage:
            prerequisite_review.append({"client_id": client_id, "reason": "not_in_current_client_staging"})
            continue
        phones = normalized_phones(stage.get("phones", ""))
        if not phones:
            prerequisite_review.append({"client_id": client_id, "reason": "invalid_phone"})
            continue
        if any(prior["client_id"] != client_id for phone in phones for prior in old_by_phone.get(phone, [])):
            prerequisite_review.append({"client_id": client_id, "reason": "phone_collision_with_old_client"})
            continue
        if phones & seen_phones.keys():
            prerequisite_review.append({"client_id": client_id, "reason": "phone_collision_with_delta_client"})
            continue
        base.assign_managers([stage], managers_by_club)
        base.assign_branches([stage], branches_by_club)
        labeled = funnel.apply_fitbase_labels([stage], funnel.CUSTOMER_SINGLE_STAGE_MODE)[0]
        if client_id in current_by_id:
            row = {h: val(current_by_id[client_id].get(h)) for h in base.MAIN_HEADERS}
        else:
            row = {"client_id": client_id, "phone": stage["phones"], "client_fio": stage["client_fio"], "email": stage["email"], "funnel": labeled["funnel"], "funnel_step": labeled["funnel_step"], "budget": 0, "create_date": stage["create_date"], "manager": stage["manager"], "филиал": stage["branch"]}
        lead_rows.append(row)
        prerequisite_rows.append({"client_ref": stage["client_ref"], "reason": "buyer_prerequisite", **row})
        known_target_ids.add(client_id)
        seen_phones.update({phone: stage for phone in phones})
    write_csv(reports / "buyer_prerequisite_leads.csv", prerequisite_rows, ["client_ref", "reason", *base.MAIN_HEADERS])
    write_csv(reports / "buyer_prerequisite_review.csv", prerequisite_review, ["client_id", "reason"])
    member_importable = []
    member_review = []
    for row in member_rows:
        if row["client_id"] in known_target_ids:
            member_importable.append(row)
        else:
            member_review.append(row)
            fact = next(f for f in selected_members if f["subscription_ref"] == row["_subscription_ref"])
            member_exclusions.append({"subscription_ref": row["_subscription_ref"], "sale_datetime": fact["sale_datetime"], "client_id": row["client_id"], "reason": "target_client_unresolved"})
    member_rows = member_importable
    if len(member_rows) + len(member_exclusions) != len(selected_members):
        raise RuntimeError("Membership manifest does not reconcile")
    purchase_row_count = len(member_rows)
    service_output = []
    configured_service_output = []
    extra_service_output = []
    service_exclusions = []
    service_review = []
    for fact, row in zip(window_services, service_rows, strict=True):
        if val(row["service_id"]) in old_service_ids or service_signature(row) in old_signatures:
            service_exclusions.append({"sale_doc_ref": fact["sale_doc_ref"], "sale_line_no": fact["sale_line_no"], "reason": "exact_row_previously_published"})
        elif not normalized_phones(row.get("phone") or "") or "розничный клиент" in (row.get("client_fio") or "").lower():
            service_exclusions.append({"sale_doc_ref": fact["sale_doc_ref"], "sale_line_no": fact["sale_line_no"], "reason": "missing_valid_phone_or_retail_placeholder"})
            service_review.append({**row, "_sale_line_no": fact["sale_line_no"], "_review_reason": "missing_valid_phone_or_retail_placeholder"})
        elif row["client_id"] not in known_target_ids:
            service_exclusions.append({"sale_doc_ref": fact["sale_doc_ref"], "sale_line_no": fact["sale_line_no"], "reason": "target_client_not_in_leads_delivery"})
            service_review.append({**row, "_sale_line_no": fact["sale_line_no"], "_review_reason": "target_client_not_in_leads_delivery"})
        else:
            row["_sale_line_no"] = fact["sale_line_no"]
            service_output.append(row)
            if (fact["sale_doc_ref"], fact["sale_line_no"]) in extra_sale_keys:
                extra_service_output.append(row)
            else:
                configured_service_output.append(row)
    if len(service_output) + len(service_exclusions) != len(window_services):
        raise RuntimeError("Service manifest does not reconcile")
    if (len(configured_window), len(extra_window), len(configured_service_output), len(extra_service_output)) != (31, 21, 30, 21):
        raise RuntimeError("Three-day service split counts changed")
    if len({val(r["service_id"]) for r in service_output}) != len(service_output):
        raise RuntimeError("Duplicate service_id across the two output files")
    candidate_lead_ids = {r["client_id"] for r in lead_rows}
    if len(lead_rows) != 72 or len(candidate_lead_ids) != 72 or prerequisite_rows:
        raise RuntimeError("Expected 72 distinct new client candidates and no prerequisite leads")
    purchase_ids = {r["client_id"] for r in member_rows}
    new_lead_ids = candidate_lead_ids & purchase_ids
    refuser_ids = existing_target_ids & purchase_ids
    if len(new_lead_ids) != 50 or len(refuser_ids) != 25:
        raise RuntimeError("Expected 50 new and 25 prior-refuser membership buyers")
    if candidate_lead_ids & existing_target_ids:
        raise RuntimeError("A prior refuser was added to the new client candidates")

    owner_change_by_new_id = defaultdict(list)
    for fact in member_facts:
        if fact.get("client_id") in candidate_lead_ids and fact.get("owner_change_ref"):
            owner_change_by_new_id[fact["client_id"]].append(fact)
    service_buyer_ids = {r["client_id"] for r in service_output}
    candidate_by_id = {r["client_id"]: r for r in lead_rows}
    excluded_new_leads = []
    owner_transfer_audit = []
    for client_id in sorted(candidate_lead_ids - new_lead_ids):
        lead = candidate_by_id[client_id]
        stage = current_stage_by_id[client_id]
        transfers = [
            fact for fact in owner_change_by_new_id.get(client_id, [])
            if val(fact.get("document_number")) in old_contract_ids
            and fact.get("sale_datetime", "") <= OLD_CUTOFF
        ]
        if transfers:
            reason = "prior_contract_transferred_to_new_owner_no_new_sale"
        elif client_id in service_buyer_ids:
            reason = "service_purchase_no_new_membership_sale"
        else:
            reason = "no_purchase"
        excluded_new_leads.append({
            "client_ref": stage["client_ref"], "client_id": client_id,
            "client_fio": lead["client_fio"], "phone": lead["phone"],
            "funnel": lead["funnel"], "funnel_step": lead["funnel_step"],
            "reason": reason, "service_purchase_in_window": "1" if client_id in service_buyer_ids else "0",
            "prior_owner_client_ids": ", ".join(sorted({val(f.get("original_client_id")) for f in transfers})),
            "prior_contract_ids": ", ".join(sorted({val(f.get("document_number")) for f in transfers})),
            "prior_subscription_refs": ", ".join(sorted({val(f.get("subscription_ref")) for f in transfers})),
            "owner_change_datetimes": ", ".join(sorted({val(f.get("owner_change_datetime")) for f in transfers})),
        })
        for fact in transfers:
            old_id = val(fact.get("original_client_id"))
            prior = old_by_id.get(old_id, [])
            if len(prior) != 1:
                raise RuntimeError(f"Prior owner is not unique in baseline: {old_id}")
            owner_transfer_audit.append({
                "old_client_ref": prior[0]["client_ref"],
                "old_client_id": old_id,
                "old_client_fio": prior[0]["client_fio"],
                "new_client_ref": stage["client_ref"],
                "new_client_id": client_id,
                "new_client_fio": lead["client_fio"],
                "contract_id": val(fact.get("document_number")),
                "subscription_ref": val(fact.get("subscription_ref")),
                "transfer_datetime": val(fact.get("owner_change_datetime")),
                "new_lead_funnel": lead["funnel"],
                "new_lead_funnel_step": lead["funnel_step"],
            })
    exclusion_counts = Counter(r["reason"] for r in excluded_new_leads)
    if exclusion_counts != {
        "prior_contract_transferred_to_new_owner_no_new_sale": 2,
        "service_purchase_no_new_membership_sale": 5,
        "no_purchase": 15,
    }:
        raise RuntimeError(f"New client non-membership cohorts changed: {exclusion_counts}")
    if len(owner_transfer_audit) != 2 or len({r["new_client_id"] for r in owner_transfer_audit}) != 2:
        raise RuntimeError("Expected two distinct prior-contract owner transfers")
    write_csv(reports / "new_lead_exclusions.csv", excluded_new_leads, [
        "client_ref", "client_id", "client_fio", "phone", "funnel", "funnel_step",
        "reason", "service_purchase_in_window", "prior_owner_client_ids",
        "prior_contract_ids", "prior_subscription_refs", "owner_change_datetimes",
    ])
    write_csv(reports / "new_client_prior_contract_transfers.csv", owner_transfer_audit, [
        "old_client_ref", "old_client_id", "old_client_fio",
        "new_client_ref", "new_client_id", "new_client_fio",
        "contract_id", "subscription_ref", "transfer_datetime",
        "new_lead_funnel", "new_lead_funnel_step",
    ])
    service_only_ids = {r["client_id"] for r in excluded_new_leads if r["reason"] == "service_purchase_no_new_membership_sale"}
    katya_service_only_output = [r for r in service_output if r["client_id"] in service_only_ids]
    configured_service_output = [r for r in configured_service_output if r["client_id"] not in service_only_ids]
    extra_service_output = [r for r in extra_service_output if r["client_id"] not in service_only_ids]
    service_partition = configured_service_output + extra_service_output + katya_service_only_output
    if (len(configured_service_output), len(extra_service_output), len(katya_service_only_output)) != (26, 20, 5):
        raise RuntimeError("Expected disjoint 26 + 20 + 5 service split")
    if {r["service_id"] for r in service_partition} != {r["service_id"] for r in service_output} or len(service_partition) != len(service_output):
        raise RuntimeError("Three service files do not reconcile to the 51 purchase rows")
    if {r["client_id"] for r in katya_service_only_output} != service_only_ids:
        raise RuntimeError("Katya service-only client IDs differ from the five excluded new leads")
    for row in katya_service_only_output:
        if Decimal(val(row["amount_of_payment"])) != Decimal(val(row["price"])) or Decimal(val(row["payment_left"])) != 0:
            raise RuntimeError("Katya service payment fields differ from the approved rule")
    if sum(r["service_name"] == "Гостевой визит" and Decimal(val(r["price"])) == 0 for r in katya_service_only_output) != 1:
        raise RuntimeError("The zero-price guest visit is missing from Katya service-only file")
    write_csv(reports / "prior_refusers_without_new_membership.csv", [
        {"client_id": client_id, "reason": "service_purchase_no_new_membership_sale" if client_id in service_buyer_ids else "no_new_membership_sale"}
        for client_id in sorted(existing_target_ids - refuser_ids)
    ], ["client_id", "reason"])
    if len(existing_target_ids - refuser_ids) != 6:
        raise RuntimeError("Expected six prior refusers without a new membership")

    lead_rows = [r for r in lead_rows if r["client_id"] in new_lead_ids]
    new_leads = [r for r in new_leads if r["client_id"] in new_lead_ids]
    if len(lead_rows) != 50 or len(new_leads) != 50:
        raise RuntimeError("New client lead rows do not reconcile with membership purchases")
    write_csv(reports / "new_leads_manifest.csv", [
        {"client_ref": r["client_ref"], **v} for r, v in zip(new_leads, lead_rows, strict=True)
    ], ["client_ref", *base.MAIN_HEADERS])
    exclusion_by_id = {r["client_id"]: r["reason"] for r in excluded_new_leads}
    for row in lead_scope:
        if row["disposition"] == "import" and row["client_id"] in exclusion_by_id:
            row["disposition"] = "exclude_" + exclusion_by_id[row["client_id"]]
    write_csv(reports / "new_lead_scope_audit.csv", lead_scope, [
        "client_ref", "client_id", "funnel", "in_current_full_export", "disposition",
    ])

    for row in member_rows:
        row["_row_kind"] = "new_purchase"
        if row["client_id"] in refuser_ids:
            row["tag"] = "отказники"
    if len(member_rows) != purchase_row_count or len(member_rows) != 213:
        raise RuntimeError("Expected 213 real membership purchase rows")
    contract_ids = [val(r["contract_id"]) for r in member_rows]
    if any(not contract_id for contract_id in contract_ids) or len(contract_ids) != len(set(contract_ids)):
        raise RuntimeError("Membership purchase contracts are blank or repeated")
    if sum(r["client_id"] in new_lead_ids for r in member_rows) != 50:
        raise RuntimeError("New lead purchases do not reconcile")
    if sum(r["client_id"] in refuser_ids for r in member_rows) != 28:
        raise RuntimeError("Prior-refuser purchases do not reconcile")
    other_old_rows = [r for r in member_rows if r["client_id"] not in new_lead_ids | refuser_ids]
    if len(other_old_rows) != 135 or len({r["client_id"] for r in other_old_rows}) != 134:
        raise RuntimeError("Other prior-buyer purchases do not reconcile")
    if any(r["tag"] for r in member_rows if r["client_id"] in new_lead_ids):
        raise RuntimeError("A new lead inherited the prior-refuser tag")
    if any(r["tag"] != "отказники" for r in member_rows if r["client_id"] in refuser_ids):
        raise RuntimeError("A prior refuser lacks the requested tag")
    if any(r["client_id"] in existing_target_ids - refuser_ids for r in member_rows):
        raise RuntimeError("A non-purchasing prior refuser was added to memberships")
    if {r["client_id"] for r in lead_rows} != new_lead_ids:
        raise RuntimeError("Lead and membership buyer sets differ")
    lead_membership_audit = []
    for lead in lead_rows:
        linked = [r for r in member_rows if r["client_id"] == lead["client_id"]]
        lead_membership_audit.append({
            "client_id": lead["client_id"],
            "client_fio": lead["client_fio"],
            "funnel": lead["funnel"],
            "funnel_step": lead["funnel_step"],
            "membership_rows": len(linked),
            "new_purchase_rows": sum(bool(val(r["contract_id"])) for r in linked),
        })
    write_csv(reports / "new_leads_membership_coverage.csv", lead_membership_audit, [
        "client_id", "client_fio", "funnel", "funnel_step", "membership_rows", "new_purchase_rows",
    ])
    write_csv(reports / "membership_import_manifest.csv", member_rows, [*memberships.CLIENT_HEADERS, "_subscription_ref", "_money_source", "_payment_match_source", "_row_kind"])
    write_csv(reports / "membership_exclusions.csv", member_exclusions, ["subscription_ref", "sale_datetime", "client_id", "reason"])
    write_csv(reports / "membership_target_client_review.csv", member_review, [*memberships.CLIENT_HEADERS, "_subscription_ref"])
    write_csv(reports / "membership_builder_issues.csv", member_issues, ["issue_type", "contract_id", "client_id", "client_fio", "contract_name", "details"])
    write_csv(reports / "service_import_manifest.csv", service_output, [*services.CLIENT_HEADERS, "_sale_doc_ref", "_sale_line_no", "_sale_datetime", "_date_state"])
    write_csv(reports / "service_import_manifest_prior_catalogue.csv", configured_service_output, [*services.CLIENT_HEADERS, "_sale_doc_ref", "_sale_line_no", "_sale_datetime", "_date_state"])
    write_csv(reports / "service_import_manifest_extra_for_katya.csv", extra_service_output, [*services.CLIENT_HEADERS, "_sale_doc_ref", "_sale_line_no", "_sale_datetime", "_date_state"])
    write_csv(reports / "service_import_manifest_5_for_katya.csv", katya_service_only_output, [*services.CLIENT_HEADERS, "_sale_doc_ref", "_sale_line_no", "_sale_datetime", "_date_state"])
    write_csv(reports / "service_exclusions.csv", service_exclusions, ["sale_doc_ref", "sale_line_no", "reason"])
    write_csv(reports / "service_target_client_review.csv", service_review, [*services.CLIENT_HEADERS, "_sale_doc_ref", "_sale_line_no", "_sale_datetime", "_review_reason"])
    write_csv(reports / "service_builder_issues.csv", service_issues, ["issue_type", "service_id", "client_id", "client_fio", "service_name", "details"])
    supplement_by_id = {r["client_id"]: r for r in supplement_clients}
    supplement_by_ref = {r["client_ref"]: r for r in supplement_clients}
    supplement_by_phone = defaultdict(list)
    for prior in supplement_clients:
        for phone in normalized_phones(prior["phones"]):
            supplement_by_phone[phone].append(prior)
    client_overlaps = []
    for row in lead_rows:
        client_id = val(row["client_id"])
        current = current_stage_by_id.get(client_id, {})
        matches = {}
        for prior in (supplement_by_id.get(client_id), supplement_by_ref.get(current.get("client_ref"))):
            if prior:
                matches[prior["client_id"]] = prior
        for prior in matches.values():
            client_overlaps.append({"delta_client_id": client_id, "delta_client_ref": current.get("client_ref", ""), "supplement_client_id": prior["client_id"], "supplement_client_ref": prior["client_ref"], "overlap": "client_id_or_ref"})
    write_csv(reports / "supplement_client_key_overlaps.csv", client_overlaps, ["delta_client_id", "delta_client_ref", "supplement_client_id", "supplement_client_ref", "overlap"])
    supplement_contract_ids = {val(r["contract_id"]) for r in supplement_members}
    supplement_service_ids = {val(r["service_id"]) for r in supplement_services}
    member_overlaps = [{"contract_id": r["contract_id"], "client_id": r["client_id"]} for r in member_rows if val(r["contract_id"]) in supplement_contract_ids]
    service_overlaps = [{"service_id": r["service_id"], "client_id": r["client_id"]} for r in service_output if val(r["service_id"]) in supplement_service_ids]
    write_csv(reports / "supplement_membership_key_overlaps.csv", member_overlaps, ["contract_id", "client_id"])
    write_csv(reports / "supplement_service_key_overlaps.csv", service_overlaps, ["service_id", "client_id"])
    shared_phones = []
    buyer_ids = {r["client_id"] for r in member_rows + service_output}
    for client_id in sorted(buyer_ids):
        current = current_stage_by_id.get(client_id)
        if not current:
            continue
        for phone in sorted(normalized_phones(current["phones"])):
            for prior in supplement_by_phone.get(phone, []):
                if prior["client_id"] != client_id:
                    shared_phones.append({"normalized_phone": phone, "delta_buyer_client_id": client_id, "supplement_client_id": prior["client_id"], "delta_buyer_fio": current["client_fio"], "supplement_client_fio": prior["client_fio"], "interpretation": "shared_phone_distinct_client_ids"})
    write_csv(reports / "supplement_shared_buyer_phones.csv", shared_phones, ["normalized_phone", "delta_buyer_client_id", "supplement_client_id", "delta_buyer_fio", "supplement_client_fio", "interpretation"])
    supplement_audit = {"baseline_status": "prepared_prior_supplement_upload_unconfirmed", "backup_finish_at": OLD_CUTOFF, "effective_cutoff_at": SUPPLEMENT_CUTOFF, "supplement_leads": len(supplement_leads), "supplement_memberships": len(supplement_members), "supplement_services": len(supplement_services), "delta_lead_client_id_or_ref_overlaps": len(client_overlaps), "delta_membership_contract_id_overlaps": len(member_overlaps), "delta_service_id_overlaps": len(service_overlaps), "shared_phone_across_distinct_buyer_and_supplement_clients": len(shared_phones)}
    (reports / "supplement_baseline_audit.json").write_text(json.dumps(supplement_audit, ensure_ascii=False, indent=2) + "\n")
    if client_overlaps or member_overlaps or service_overlaps:
        raise RuntimeError(f"Delta repeats supplement keys: {supplement_audit}")
    if any(val(r["contract_id"]) in old_contract_ids for r in member_rows if val(r["contract_id"])):
        raise RuntimeError("Delta membership contract_id repeats a prepared prior delivery")
    if any(val(r["service_id"]) in old_service_ids for r in service_output):
        raise RuntimeError("Delta service_id repeats a prepared prior delivery")
    raw_sales = csv_rows(NEW / "raw/staging/stg_sales_all.csv")
    raw_window = [r for r in raw_sales if in_window(r.get("sale_datetime"))]
    service_scope = Counter((r.get("sale_source", ""), r.get("product_class", "")) for r in raw_window)
    inventory_path = OUT / "reports/scope/all_sale_lines_by_product.txt"
    product_scope = []
    if inventory_path.exists():
        membership_names = {r.get("subscription_name", "") for r in selected_members}
        included_service_names = {r["service_name"] for r in window_services}
        retail_markers = ("БонАква", "Батончик", "Печенье", "Крем ", "Соус ", "Гель для душа", "Шампунь", "Тапочки", "Стикини", "Шапочка", "Полотенце")
        for line in inventory_path.read_text().splitlines():
            parts = line.split("|")
            if len(parts) != 6:
                raise RuntimeError(f"Malformed sale-line inventory: {line[:100]}")
            ref, linked_type, code, name, count, amount = parts
            if name in included_service_names:
                classification = "service_included"
            elif name in membership_names and linked_type == "000000A3":
                classification = "membership_included"
            elif name.startswith(retail_markers):
                classification = "retail_excluded"
            elif linked_type == "0000008A":
                classification = "modifier_or_administrative_review"
            else:
                classification = "manual_classification_review"
            product_scope.append({"product_ref": ref, "linked_type": linked_type, "product_code": code, "product_name": name, "sale_lines": count, "line_amount_sum": amount, "classification": classification})
        write_csv(reports / "all_window_sale_products_classification.csv", product_scope, ["product_ref", "linked_type", "product_code", "product_name", "sale_lines", "line_amount_sum", "classification"])
        if sum(int(r["sale_lines"]) for r in product_scope) != 425:
            raise RuntimeError("Sale-line inventory total changed")
    summary = {
        "old_backup_finish_at": OLD_CUTOFF,
        "new_backup_finish_at": NEW_CUTOFF,
        "raw_new_client_refs": len(raw_new_leads),
        "new_client_candidate_rows": len(candidate_lead_ids),
        "new_client_lead_rows": len(new_leads),
        "new_client_exclusions": dict(exclusion_counts),
        "existing_refuser_buyer_cards_confirmed": len(existing_target_ids),
        "existing_refuser_buyer_cards_unresolved": len(existing_audit) - len(existing_target_ids),
        "existing_refuser_fitbase_snapshot_at_utc": snapshot_at,
        "prior_refusers_without_new_membership": len(existing_target_ids - refuser_ids),
        "buyer_prerequisite_lead_rows": len(prerequisite_rows),
        "buyer_prerequisite_unresolved_clients": len(prerequisite_review),
        "lead_rows": len(lead_rows),
        "lead_collisions": len(collisions),
        "membership_window_facts": len(selected_members),
        "membership_rows": len(member_rows),
        "membership_purchase_rows": purchase_row_count,
        "membership_new_lead_clients": len(new_lead_ids),
        "membership_prior_refuser_clients": len(refuser_ids),
        "membership_prior_refuser_purchase_rows": sum(r["client_id"] in refuser_ids for r in member_rows),
        "membership_other_prior_buyer_clients": len({r["client_id"] for r in other_old_rows}),
        "membership_other_prior_buyer_rows": len(other_old_rows),
        "membership_exclusions": len(member_exclusions),
        "service_window_facts_configured": len(configured_window),
        "service_window_facts_extra": len(extra_window),
        "service_rows": len(service_output),
        "service_rows_prior_catalogue": len(configured_service_output),
        "service_rows_extra_for_katya": len(extra_service_output),
        "service_rows_5_for_katya": len(katya_service_only_output),
        "service_extra_names": sorted(extra_names),
        "service_exclusions": len(service_exclusions),
        "service_names_in_window": service_names,
        "all_posted_sale_lines_window": sum(int(r["sale_lines"]) for r in product_scope),
        "product_scope_classification": dict(Counter(r["classification"] for r in product_scope)),
        "raw_sales_window_by_source_class": {f"{source}|{kind}": count for (source, kind), count in service_scope.items()},
        "service_payment_rule": "Customer confirmed amount_of_payment=price and payment_left=0 for service imports.",
        "service_scope_limitation": "Clearly service sale lines are mapped subject to target-client checks. Other sale products are categorized in all_window_sale_products_classification.csv; ambiguous lines need business classification before claiming all services.",
    }
    (reports / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n")
    workbook_data = [
        ("fitbase_delta_import_zayavki_20260923.xlsx", base.MAIN_HEADERS, base.MAIN_RUS_HEADERS, lead_rows),
        ("fitbase_delta_import_abonementy_clientov_20260923.xlsx", memberships.CLIENT_HEADERS, memberships.CLIENT_RUS_HEADERS, member_rows),
        ("fitbase_delta_import_uslugi_clientov_20260923.xlsx", services.CLIENT_HEADERS, services.CLIENT_RUS_HEADERS, configured_service_output),
        ("Дополнительные_услуги_20260923.xlsx", services.CLIENT_HEADERS, services.CLIENT_RUS_HEADERS, extra_service_output),
        ("Услуги_без_абонемента_20260923.xlsx", services.CLIENT_HEADERS, services.CLIENT_RUS_HEADERS, katya_service_only_output),
    ]
    specs = []
    for name, headers, rus, rows in workbook_data:
        specs.append({"name": name, "headers": headers, "russian_headers": rus, "rows": [[val(r.get(h)) for h in headers] for r in rows]})
    (reports / "workbook_specs.json").write_text(json.dumps(specs, ensure_ascii=False))
    to_write = [name for name, headers, _, rows in workbook_data if not workbook_matches(OUT / name, headers, rows)]
    if to_write:
        subprocess.run(["node", str(ROOT / "scripts/write_20260923_delta_imports.mjs"), *to_write], check=True)
    (reports / "membership_client_only_rows.csv").unlink(missing_ok=True)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.parse_args()
    build()
