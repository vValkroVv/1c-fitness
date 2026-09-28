#!/usr/bin/env python3
"""Build a source-database photo archive for the exact 75 Sep 23 delta buyers.

The two import manifests define the client groups and their delivered phones.
The restored Sep 23 database supplies photo bytes. Refusers remain explicitly
tagged as prior_refuser_new_membership_buyer; their funnel is not inferred.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import io
import json
import logging
import re
import subprocess
import sys
import textwrap
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DELIVERY = ROOT / "output/20260923_delta_from_20260920"
IMPORT_REPORTS = DELIVERY / "reports/imports"
LEADS_MANIFEST = IMPORT_REPORTS / "new_leads_manifest.csv"
MEMBERSHIP_MANIFEST = IMPORT_REPORTS / "membership_import_manifest.csv"
MEMBERSHIP_XLSX = DELIVERY / "fitbase_delta_import_abonementy_clientov_20260923.xlsx"

DEFAULT_CONTAINER = "mssql-fitness-2022"
DEFAULT_DATABASE = "FitnessRestored_20260923_macos"
DEFAULT_BACKUP_UUID = "98038663-5828-4efe-80e4-327dbff24e3e"
DEFAULT_BACKUP_FINISH_AT = "2026-09-23 23:36:39"
EXPECTED_LEADS = 50
EXPECTED_PRIOR_REFUSER_BUYERS = 25
EXPECTED_TARGETS = 75

PHOTO_GROUP = "new_membership_buyer"
REFUSER_GROUP = "prior_refuser_new_membership_buyer"
REFUSER_TAG = "отказники"
SAFE_SQL_IDENTIFIER = re.compile(r"^[A-Za-z0-9_]+$")
SAFE_CONTAINER = re.compile(r"^[A-Za-z0-9_.-]+$")
SAFE_CLIENT_ID = re.compile(r"^[0-9]{9}$")


@dataclass(frozen=True)
class Target:
    client_id: str
    group: str
    client_fio: str
    phone: str
    normalized_phones: tuple[str, ...]


@dataclass(frozen=True)
class Candidate:
    client_id: str
    raw_phone: str
    source: str
    metadata_extension: str

    @property
    def key(self) -> tuple[str, str, str, str]:
        return self.client_id, self.raw_phone, self.source, self.metadata_extension


@dataclass(frozen=True)
class UsablePhoto:
    candidate: Candidate
    detected_extension: str
    jpeg: bytes


def import_photo_helpers():
    path = ROOT / "scripts/41_export_active_client_photos_zip.py"
    spec = importlib.util.spec_from_file_location("photo_export_helpers", path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Could not load photo export helpers from {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def csv_read(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sort_id(client_id: str) -> tuple[int, str]:
    return int(client_id), client_id


def read_targets(photo: Any) -> tuple[dict[str, Target], dict[str, int]]:
    """Derive the two exact groups from the customer-facing import manifests."""

    leads = csv_read(LEADS_MANIFEST)
    memberships = csv_read(MEMBERSHIP_MANIFEST)
    required_lead_fields = {"client_id", "phone", "client_fio"}
    required_membership_fields = {
        "tag", "contract_id", "client_id", "phone", "client_fio",
        "_subscription_ref", "_row_kind",
    }
    if not leads or not required_lead_fields.issubset(leads[0]):
        raise ValueError(f"Lead manifest is empty or missing fields: {LEADS_MANIFEST}")
    if not memberships or not required_membership_fields.issubset(memberships[0]):
        raise ValueError(f"Membership manifest is empty or missing fields: {MEMBERSHIP_MANIFEST}")

    lead_ids: set[str] = set()
    targets: dict[str, Target] = {}
    for row in leads:
        client_id = row["client_id"].strip()
        if not SAFE_CLIENT_ID.fullmatch(client_id):
            raise ValueError(f"Unsafe or missing client_id in lead manifest: {client_id!r}")
        if client_id in lead_ids:
            raise ValueError(f"Duplicate lead client_id in manifest: {client_id}")
        lead_ids.add(client_id)
        phone = row["phone"].strip()
        targets[client_id] = Target(
            client_id, PHOTO_GROUP, row["client_fio"].strip(), phone,
            photo.normalize_phones(phone),
        )
    if len(leads) != EXPECTED_LEADS or len(lead_ids) != EXPECTED_LEADS:
        raise ValueError(f"Expected {EXPECTED_LEADS} distinct new leads; found rows={len(leads)}, IDs={len(lead_ids)}")

    tagged_rows = [row for row in memberships if row["tag"].strip() == REFUSER_TAG]
    qualified_refuser_rows = [
        row for row in tagged_rows
        if row["_row_kind"].strip() == "new_purchase"
        and row["contract_id"].strip()
        and row["_subscription_ref"].strip()
    ]
    tagged_ids = {row["client_id"].strip() for row in tagged_rows}
    qualified_ids = {row["client_id"].strip() for row in qualified_refuser_rows}
    if tagged_ids != qualified_ids:
        raise ValueError("Some tagged отказники have no qualifying real new contract")
    if not tagged_rows or len(qualified_ids) != EXPECTED_PRIOR_REFUSER_BUYERS:
        raise ValueError(
            f"Expected {EXPECTED_PRIOR_REFUSER_BUYERS} distinct tagged refusal buyers; "
            f"found tagged rows={len(tagged_rows)}, IDs={len(qualified_ids)}"
        )

    by_refuser: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in qualified_refuser_rows:
        client_id = row["client_id"].strip()
        if not SAFE_CLIENT_ID.fullmatch(client_id):
            raise ValueError(f"Unsafe or missing refuser client_id: {client_id!r}")
        by_refuser[client_id].append(row)
    for client_id, rows in by_refuser.items():
        if client_id in lead_ids:
            raise ValueError(f"Client {client_id} overlaps new_leads and отказники groups")
        phones = {row["phone"].strip() for row in rows}
        fios = {row["client_fio"].strip() for row in rows if row["client_fio"].strip()}
        if len(phones) != 1 or len(fios) > 1:
            raise ValueError(f"Conflicting phone/FIO across tagged contracts for client {client_id}")
        row = sorted(rows, key=lambda item: (item["contract_id"], item["_subscription_ref"]))[0]
        phone = row["phone"].strip()
        targets[client_id] = Target(
            client_id, REFUSER_GROUP, next(iter(fios), ""), phone,
            photo.normalize_phones(phone),
        )

    if len(targets) != EXPECTED_TARGETS:
        raise ValueError(f"Expected {EXPECTED_TARGETS} non-overlapping targets; found {len(targets)}")
    if len(lead_ids & qualified_ids) != 0:
        raise ValueError("Lead and refusal buyer IDs overlap")

    counts = {
        "new_lead_manifest_rows": len(leads),
        "new_lead_clients": len(lead_ids),
        "tagged_refuser_contract_rows": len(tagged_rows),
        "qualified_refuser_contract_rows": len(qualified_refuser_rows),
        "prior_refuser_new_membership_buyers": len(qualified_ids),
        "target_clients": len(targets),
    }
    return targets, counts


def validate_final_membership(targets: dict[str, Target]) -> dict[str, int]:
    """Check the published membership workbook, not only its source manifest."""

    manifest = csv_read(MEMBERSHIP_MANIFEST)
    workbook = load_workbook(MEMBERSHIP_XLSX, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        stream = sheet.iter_rows(values_only=True)
        machine_header = next(stream, None)
        russian_header = next(stream, None)
        if machine_header is None or russian_header is None:
            raise ValueError(f"Membership workbook is missing its two header rows: {MEMBERSHIP_XLSX}")
        indexes = {str(value).strip(): idx for idx, value in enumerate(machine_header) if value is not None}
        if not {"contract_id", "client_id"}.issubset(indexes):
            raise ValueError("Membership workbook lacks contract_id or client_id machine headers")
        data_rows: list[tuple[str, str]] = []
        for row in stream:
            if not any(value is not None for value in row):
                continue
            contract_id = str(row[indexes["contract_id"]]).strip() if row[indexes["contract_id"]] is not None else ""
            client_id = str(row[indexes["client_id"]]).strip() if row[indexes["client_id"]] is not None else ""
            if contract_id or client_id:
                data_rows.append((contract_id, client_id))
    finally:
        workbook.close()

    manifest_pairs = {(row["contract_id"].strip(), row["client_id"].strip()) for row in manifest}
    workbook_pairs = set(data_rows)
    if len(data_rows) != len(manifest) or workbook_pairs != manifest_pairs:
        raise ValueError(
            f"Membership workbook rows do not match manifest: "
            f"xlsx_rows={len(data_rows)}, manifest_rows={len(manifest)}, "
            f"missing={len(manifest_pairs - workbook_pairs)}, extra={len(workbook_pairs - manifest_pairs)}"
        )
    target_ids = set(targets)
    final_ids = {client_id for _, client_id in data_rows}
    missing = sorted(target_ids - final_ids, key=sort_id)
    if missing:
        raise ValueError(f"Targets missing from final membership workbook: {missing}")
    return {
        "membership_manifest_rows": len(manifest),
        "membership_workbook_rows": len(data_rows),
        "target_ids_in_final_membership_rows": len(target_ids & final_ids),
    }


def validate_args(args: argparse.Namespace, photo: Any) -> tuple[str, str]:
    if not SAFE_SQL_IDENTIFIER.fullmatch(args.database):
        raise ValueError(f"Unsafe database name: {args.database!r}")
    if not SAFE_CONTAINER.fullmatch(args.container):
        raise ValueError(f"Unsafe Docker container name: {args.container!r}")
    cutoff = datetime.strptime(args.cutoff_at, "%Y-%m-%d %H:%M:%S")
    backup = datetime.strptime(args.backup_finish_at, "%Y-%m-%d %H:%M:%S")
    if cutoff.strftime("%Y-%m-%d %H:%M:%S") != args.cutoff_at:
        raise ValueError("cutoff_at must be a canonical YYYY-MM-DD HH:MM:SS timestamp")
    if backup.strftime("%Y-%m-%d %H:%M:%S") != args.backup_finish_at:
        raise ValueError("backup_finish_at must be a canonical YYYY-MM-DD HH:MM:SS timestamp")
    try:
        import uuid
        args.backup_set_uuid = str(uuid.UUID(args.backup_set_uuid))
    except ValueError as exc:
        raise ValueError(f"Invalid backup set UUID: {args.backup_set_uuid}") from exc
    return cutoff.strftime("%Y%m%d"), backup.strftime("%Y-%m-%d %H:%M:%S")


def verify_restore_identity(args: argparse.Namespace) -> dict[str, str]:
    """Read the newest live restorehistory row before any photo extraction."""

    sql = f"""SET NOCOUNT ON;
SELECT TOP (1) CONVERT(varchar(36), bs.backup_set_uuid) AS backup_set_uuid,
       CONVERT(varchar(19), bs.backup_finish_date, 120) AS backup_finish_at,
       CONVERT(varchar(33), rh.restore_date, 126) AS restore_date,
       rh.destination_database_name AS destination_database_name
FROM msdb.dbo.restorehistory AS rh
JOIN msdb.dbo.backupset AS bs ON bs.backup_set_id = rh.backup_set_id
WHERE rh.destination_database_name = N'{args.database}' AND rh.restore_type = 'D'
ORDER BY rh.restore_date DESC
FOR JSON PATH;"""
    command = [
        "docker", "exec", args.container, "/bin/bash", "-lc",
        'exec /opt/mssql-tools18/bin/sqlcmd -S localhost -U sa -P "$MSSQL_SA_PASSWORD" -C "$@"',
        "sqlcmd", "-d", args.database, "-b", "-h", "-1", "-Q", sql, "-w", "65535",
    ]
    result = subprocess.run(command, cwd=ROOT, text=True, capture_output=True)
    if result.returncode != 0:
        raise RuntimeError(f"Restore identity query failed: {result.stderr.strip() or result.stdout.strip()}")
    output = result.stdout.strip()
    start, end = output.find("[{"), output.rfind("}]")
    if start < 0 or end < start:
        raise RuntimeError(f"Could not parse restorehistory JSON: {output[:1000]}")
    rows = json.loads(output[start : end + 2])
    if len(rows) != 1:
        raise RuntimeError(f"Expected one latest database restore row, got {len(rows)}")
    actual = rows[0]
    expected_uuid = args.backup_set_uuid.lower()
    if str(actual["backup_set_uuid"]).lower() != expected_uuid:
        raise RuntimeError(
            f"Restored backup UUID mismatch: expected {expected_uuid}, got {actual['backup_set_uuid']}"
        )
    if actual["backup_finish_at"] != args.backup_finish_at:
        raise RuntimeError(
            f"Restored backup finish mismatch: expected {args.backup_finish_at}, got {actual['backup_finish_at']}"
        )
    if actual["destination_database_name"] != args.database:
        raise RuntimeError(f"Unexpected restored database: {actual['destination_database_name']}")
    return actual


def selected_photo_query(photo: Any, args: argparse.Namespace, ids: set[str], *, include_blob: bool, cutoff_date: str) -> str:
    if not ids:
        raise ValueError("No targets with valid phones; refusing to issue an empty photo query")
    if any(not SAFE_CLIENT_ID.fullmatch(client_id) for client_id in ids):
        raise ValueError("Unsafe client ID in source photo query")
    query = photo.sql_selected_union(args.database, include_blob=include_blob, cutoff_date=cutoff_date)
    join = (
        f"JOIN [{args.database}].fitbase_part2.final_funnel_clients AS a "
        "ON a.client_ref = CONVERT(varchar(32), c._IDRRef, 2) "
    )
    condition = f"a.cutoff_date = CONVERT(date, '{cutoff_date}', 112)"
    if query.count(join) != 2 or query.count(condition) != 2:
        raise RuntimeError("Expected photo-query funnel joins/cutoff clauses were not found exactly twice")
    id_list = ",".join(f"'{client_id}'" for client_id in sorted(ids, key=sort_id))
    query = query.replace(join, "")
    query = query.replace(condition, f"c._Code IN ({id_list})")
    if "final_funnel_clients" in query or "a.cutoff_date" in query:
        raise RuntimeError("Photo query was not fully restricted to explicit target client IDs")
    return query


def bcp_rows(photo: Any, container: str, query: str, *, include_blob: bool) -> list[list[str]]:
    """Collect the exact small client subset and fail on malformed BCP output."""

    process = photo.start_bcp(container, query)
    rows: list[list[str]] = []
    unexpected: list[bytes] = []
    expected_columns = 5 if include_blob else 4
    try:
        assert process.stdout is not None
        for raw_line in process.stdout:
            line = raw_line.rstrip(b"\r\n")
            fields = line.split(b"\t", expected_columns - 1)
            if len(fields) != expected_columns:
                if not photo.is_bcp_status_line(line):
                    unexpected.append(line[:500])
                continue
            rows.append([field.decode("utf-8", errors="strict") for field in fields])
        photo.finish_bcp(process, unexpected)
    finally:
        photo.close_bcp(process)
    return rows


def parse_candidates(rows: list[list[str]], targets: dict[str, Target]) -> dict[str, list[Candidate]]:
    candidates: dict[str, dict[tuple[str, str, str, str], Candidate]] = defaultdict(dict)
    allowed_extensions = {"jpg", "jpeg", "png", "bmp", "gif"}
    for client_id, raw_phone, source, extension in rows:
        if client_id not in targets:
            raise RuntimeError(f"Source photo query returned out-of-scope client {client_id}")
        if source not in {"Reference65_main", "Reference115_copy"}:
            raise RuntimeError(f"Unexpected source photo kind for {client_id}: {source!r}")
        extension = extension.lower()
        if extension not in allowed_extensions:
            raise RuntimeError(f"Unexpected source photo extension for {client_id}: {extension!r}")
        candidate = Candidate(client_id, raw_phone, source, extension)
        if candidate.key in candidates[client_id]:
            raise RuntimeError(f"Duplicate source photo metadata row: {candidate.key}")
        candidates[client_id][candidate.key] = candidate
    return {
        client_id: sorted(values.values(), key=lambda item: (item.source != "Reference65_main", item.source, item.metadata_extension))
        for client_id, values in candidates.items()
    }


def assign_photo_names(targets: dict[str, Target]) -> tuple[dict[str, str], Counter[str]]:
    owners: dict[str, set[str]] = defaultdict(set)
    for client_id, target in targets.items():
        if target.normalized_phones:
            owners[target.normalized_phones[0]].add(client_id)
    names: dict[str, str] = {}
    for client_id, target in targets.items():
        if not target.normalized_phones:
            continue
        phone = target.normalized_phones[0]
        suffix = f"__{client_id}" if len(owners[phone]) > 1 else ""
        names[client_id] = f"{phone}{suffix}.jpg"
    return names, Counter({phone: len(ids) for phone, ids in owners.items() if len(ids) > 1})


def csv_bytes(rows: list[dict[str, Any]], fieldnames: list[str]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=fieldnames, extrasaction="ignore", lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")


def write_zip(
    *, photo: Any, output: Path, inner_dir: str, cutoff_at: str,
    targets: dict[str, Target], photos: dict[str, UsablePhoto], filenames: dict[str, str],
    missing_rows: list[dict[str, str]],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    manifest: list[dict[str, Any]] = []
    for client_id in sorted(photos, key=sort_id):
        target = targets[client_id]
        usable = photos[client_id]
        manifest.append({
            "group": target.group,
            "client_id": client_id,
            "client_fio": target.client_fio,
            "phone": target.phone,
            "normalized_phone": target.normalized_phones[0],
            "photo_name": filenames[client_id],
            "sha256": hashlib.sha256(usable.jpeg).hexdigest(),
            "bytes": len(usable.jpeg),
            "photo_source": usable.candidate.source,
            "source_extension": usable.candidate.metadata_extension,
            "detected_extension": usable.detected_extension,
            "source_raw_phone": usable.candidate.raw_phone,
        })
    manifest_fields = [
        "group", "client_id", "client_fio", "phone", "normalized_phone", "photo_name", "sha256",
        "bytes", "photo_source", "source_extension", "detected_extension", "source_raw_phone",
    ]
    if len({row["client_id"] for row in manifest}) != len(manifest):
        raise RuntimeError("Archive manifest contains a client more than once")
    if len({row["photo_name"] for row in manifest}) != len(manifest):
        raise RuntimeError("Archive manifest contains duplicate JPEG names")
    output.parent.mkdir(parents=True, exist_ok=True)
    partial = output.with_name(output.name + ".partial")
    if partial.exists():
        raise FileExistsError(f"An unfinished photo archive exists; inspect before retrying: {partial}")
    try:
        with zipfile.ZipFile(partial, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6, allowZip64=True) as archive:
            archive.writestr(photo.zip_info(inner_dir, cutoff_at, directory=True), b"")
            archive.writestr(photo.zip_info(f"{inner_dir}/photos", cutoff_at, directory=True), b"")
            for client_id in sorted(photos, key=sort_id):
                archive.writestr(
                    photo.zip_info(f"{inner_dir}/photos/{filenames[client_id]}", cutoff_at),
                    photos[client_id].jpeg,
                )
            archive.writestr(
                photo.zip_info(f"{inner_dir}/_reports/manifest.csv", cutoff_at),
                csv_bytes(manifest, manifest_fields),
            )
            group_counts = Counter(row["group"] for row in manifest)
            readme = textwrap.dedent(f"""\
                Фотографии клиентов с новыми покупками из дельты за 23.09.2026.

                Фактическое завершение backup: {DEFAULT_BACKUP_FINISH_AT}
                Срез данных cutoff_at: {cutoff_at}
                Фотографий: {len(manifest)}
                Новые клиенты: {group_counts.get(PHOTO_GROUP, 0)}
                Покупатели из группы отказников: {group_counts.get(REFUSER_GROUP, 0)}

                Группа отказников не размечена как воронка: её прежняя воронка неизвестна.
                Имена JPEG: нормализованный телефон; при общем телефоне добавлен ID клиента.
                Поля группы, ID, ФИО, телефона и SHA-256 приведены в _reports/manifest.csv.
                Цели без пригодного фото и причины указаны в _reports/missing_photos.csv.
            """)
            archive.writestr(photo.zip_info(f"{inner_dir}/README.txt", cutoff_at), readme.encode("utf-8"))
            missing_fields = ["group", "client_id", "client_fio", "phone", "reason", "source_photo_errors"]
            archive.writestr(
                photo.zip_info(f"{inner_dir}/_reports/missing_photos.csv", cutoff_at),
                csv_bytes(missing_rows, missing_fields),
            )

        validation = validate_archive(
            photo, partial, inner_dir, targets, photos, filenames, missing_rows,
        )
        if output.exists() and not ARGS.overwrite:
            raise FileExistsError(f"Output appeared while exporting: {output}")
        partial.replace(output)
    except Exception:
        partial.unlink(missing_ok=True)
        raise
    return manifest, validation


def validate_archive(
    photo: Any, path: Path, inner_dir: str, targets: dict[str, Target],
    photos: dict[str, UsablePhoto], filenames: dict[str, str], missing_rows: list[dict[str, str]],
) -> dict[str, Any]:
    prefix = f"{inner_dir}/photos/"
    expected_ids = set(photos)
    with zipfile.ZipFile(path, "r") as archive:
        bad_crc = archive.testzip()
        if bad_crc is not None:
            raise RuntimeError(f"ZIP CRC validation failed at {bad_crc}")
        names = archive.namelist()
        if len(names) != len(set(names)):
            raise RuntimeError("ZIP contains duplicate entry paths")
        photo_infos = [item for item in archive.infolist() if item.filename.startswith(prefix) and not item.is_dir()]
        if len(photo_infos) != len(expected_ids):
            raise RuntimeError(f"ZIP photo count mismatch: {len(photo_infos)} != {len(expected_ids)}")
        manifest_name = f"{inner_dir}/_reports/manifest.csv"
        with archive.open(manifest_name) as handle:
            rows = list(csv.DictReader(io.TextIOWrapper(handle, encoding="utf-8-sig", newline="")))
        if len(rows) != len(expected_ids) or {row["client_id"] for row in rows} != expected_ids:
            raise RuntimeError("Archive manifest IDs do not exactly match exported target IDs")
        missing_name = f"{inner_dir}/_reports/missing_photos.csv"
        with archive.open(missing_name) as handle:
            archived_missing = list(csv.DictReader(io.TextIOWrapper(handle, encoding="utf-8-sig", newline="")))
        missing_ids = {row["client_id"] for row in archived_missing}
        expected_missing_ids = {row["client_id"] for row in missing_rows}
        if archived_missing != missing_rows or missing_ids != expected_missing_ids:
            raise RuntimeError("Archive missing-photo list differs from the external target exclusions")
        if expected_ids & missing_ids or expected_ids | missing_ids != set(targets):
            raise RuntimeError("Archive photo and missing-photo IDs do not partition all targets")
        if {row["group"] for row in rows} - {PHOTO_GROUP, REFUSER_GROUP}:
            raise RuntimeError("Archive manifest contains an unknown group")
        manifest_by_name = {row["photo_name"]: row for row in rows}
        if len(manifest_by_name) != len(rows):
            raise RuntimeError("Archive manifest contains duplicate photo names")
        archive_names = {info.filename.removeprefix(prefix) for info in photo_infos}
        if archive_names != set(manifest_by_name):
            raise RuntimeError("Archive photo names and manifest names differ")
        for info in photo_infos:
            name = info.filename.removeprefix(prefix)
            row = manifest_by_name[name]
            client_id = row["client_id"]
            target = targets[client_id]
            payload = archive.read(info)
            if row["group"] != target.group or row["client_fio"] != target.client_fio or row["phone"] != target.phone:
                raise RuntimeError(f"Manifest group/client data mismatch for {client_id}")
            if name != filenames[client_id] or row["normalized_phone"] != target.normalized_phones[0]:
                raise RuntimeError(f"Filename or normalized phone mismatch for {client_id}")
            if int(row["bytes"]) != len(payload) or row["sha256"] != hashlib.sha256(payload).hexdigest():
                raise RuntimeError(f"Photo size or SHA-256 mismatch for {client_id}")
            with Image.open(io.BytesIO(payload)) as decoded:
                if decoded.format != "JPEG":
                    raise RuntimeError(f"Non-JPEG photo in archive: {name}")
                decoded.load()
        total_bytes = sum(info.file_size for info in photo_infos)
    return {
        "zip_bytes": path.stat().st_size,
        "zip_sha256": sha256_file(path),
        "photo_entries": len(photo_infos),
        "manifest_rows": len(rows),
        "missing_photo_rows": len(missing_rows),
        "photo_uncompressed_bytes": total_bytes,
        "crc_check": "PASS",
        "unique_paths_check": "PASS",
        "decoded_jpegs_check": "PASS",
        "manifest_ids_hashes_check": "PASS",
    }


def atomic_write(path: Path, payload: bytes, *, overwrite: bool) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists() and not overwrite:
        raise FileExistsError(f"Output already exists (use --overwrite intentionally): {path}")
    partial = path.with_name(path.name + ".partial")
    if partial.exists():
        raise FileExistsError(f"An unfinished report exists; inspect before retrying: {partial}")
    partial.write_bytes(payload)
    partial.replace(path)


def make_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--container", default=DEFAULT_CONTAINER)
    parser.add_argument("--database", default=DEFAULT_DATABASE)
    parser.add_argument("--backup-set-uuid", default=DEFAULT_BACKUP_UUID)
    parser.add_argument("--backup-finish-at", default=DEFAULT_BACKUP_FINISH_AT)
    parser.add_argument("--cutoff-at", default=DEFAULT_BACKUP_FINISH_AT,
                        help="Effective cutoff as YYYY-MM-DD HH:MM:SS; defaults to backup_finish_at")
    parser.add_argument("--output", type=Path,
                        default=DELIVERY / "fitbase_client_photos_20260923_delta_75.zip")
    parser.add_argument("--overwrite", action="store_true")
    return parser


def main() -> int:
    global ARGS
    ARGS = make_parser().parse_args()
    photo = import_photo_helpers()
    cutoff_date, backup_finish_at = validate_args(ARGS, photo)

    report_dir = DELIVERY / "reports/photos"
    log_dir = DELIVERY / "logs"
    report_json_path = report_dir / "photo_archive_report.json"
    report_md_path = report_dir / "photo_archive_report.md"
    missing_path = report_dir / "missing_photos.csv"
    log_path = log_dir / "photo_export_20260923.log"
    output = ARGS.output.resolve()
    destinations = [output, report_json_path, report_md_path, missing_path, log_path]
    if not ARGS.overwrite:
        existing = [str(path) for path in destinations if path.exists()]
        if existing:
            raise FileExistsError("Outputs already exist; inspect them or pass --overwrite: " + ", ".join(existing))

    log_dir.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("fitbase_delta_photos")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    formatter = logging.Formatter("%(asctime)s %(levelname)s %(message)s")
    file_handler = logging.FileHandler(log_path, mode="w", encoding="utf-8")
    file_handler.setFormatter(formatter)
    stream_handler = logging.StreamHandler()
    stream_handler.setFormatter(formatter)
    logger.addHandler(file_handler)
    logger.addHandler(stream_handler)

    try:
        targets, target_counts = read_targets(photo)
        membership_checks = validate_final_membership(targets)
        logger.info("Derived %d targets: %s", len(targets), target_counts)
        logger.info("Verified final membership rows: %s", membership_checks)

        # Confirm the precise restored source backup immediately before any photo query.
        provenance = verify_restore_identity(ARGS)
        logger.info("Restore identity verified: %s", provenance)

        names, shared_phones = assign_photo_names(targets)
        valid_targets = {client_id for client_id, target in targets.items() if target.normalized_phones}
        no_phone = {
            client_id: ("invalid_phone" if re.search(r"\d", targets[client_id].phone) else "missing_phone")
            for client_id in targets if client_id not in valid_targets
        }
        logger.info("Valid-phone targets: %d; missing/invalid phone: %d", len(valid_targets), len(no_phone))

        metadata_query = selected_photo_query(
            photo, ARGS, valid_targets, include_blob=False, cutoff_date=cutoff_date,
        )
        candidate_rows = bcp_rows(photo, ARGS.container, metadata_query, include_blob=False)
        candidates = parse_candidates(candidate_rows, targets)
        logger.info("Source metadata candidates: %d rows for %d IDs", len(candidate_rows), len(candidates))

        blob_query = selected_photo_query(
            photo, ARGS, valid_targets, include_blob=True, cutoff_date=cutoff_date,
        )
        blob_rows = bcp_rows(photo, ARGS.container, blob_query, include_blob=True)
        expected_candidates = {candidate.key for rows in candidates.values() for candidate in rows}
        seen_candidates: set[tuple[str, str, str, str]] = set()
        usable_candidates: dict[str, list[UsablePhoto]] = defaultdict(list)
        candidate_errors: list[dict[str, str]] = []
        for client_id, raw_phone, source, extension, hex_value in blob_rows:
            candidate = Candidate(client_id, raw_phone, source, extension.lower())
            if candidate.key not in expected_candidates:
                raise RuntimeError(f"BLOB pass returned unlisted/out-of-scope photo candidate: {candidate.key}")
            if candidate.key in seen_candidates:
                raise RuntimeError(f"Duplicate source photo BLOB row: {candidate.key}")
            seen_candidates.add(candidate.key)
            try:
                raw_image, detected_extension = photo.decode_1c_photo(hex_value.encode("ascii"))
                jpeg = photo.jpeg_payload(raw_image, detected_extension)
                # Decode the final bytes now, before this candidate can be marked usable.
                with Image.open(io.BytesIO(jpeg)) as decoded:
                    if decoded.format != "JPEG":
                        raise ValueError(f"conversion produced {decoded.format}, not JPEG")
                    decoded.load()
                usable_candidates[client_id].append(UsablePhoto(candidate, detected_extension, jpeg))
            except Exception as exc:
                candidate_errors.append({
                    "client_id": client_id,
                    "photo_source": source,
                    "reason": f"{type(exc).__name__}: {exc}",
                })
                logger.warning("Unusable source photo candidate: client=%s source=%s reason=%s", client_id, source, exc)
        if seen_candidates != expected_candidates:
            missing_candidates = expected_candidates - seen_candidates
            raise RuntimeError(f"BLOB pass omitted source photo candidates: {len(missing_candidates)}")

        photos: dict[str, UsablePhoto] = {}
        for client_id, usable in usable_candidates.items():
            photos[client_id] = sorted(
                usable,
                key=lambda item: (item.candidate.source != "Reference65_main", item.candidate.source),
            )[0]

        missing_rows: list[dict[str, str]] = []
        for client_id in sorted(set(targets) - set(photos), key=sort_id):
            target = targets[client_id]
            if client_id in no_phone:
                reason = no_phone[client_id]
            elif client_id not in candidates:
                reason = "no_source_photo"
            elif not usable_candidates.get(client_id):
                reason = "no_usable_source_photo"
            else:  # Defensive; every usable candidate should have been selected above.
                reason = "source_selection_failed"
            errors = [item["reason"] for item in candidate_errors if item["client_id"] == client_id]
            missing_rows.append({
                "group": target.group, "client_id": client_id, "client_fio": target.client_fio,
                "phone": target.phone, "reason": reason,
                "source_photo_errors": " | ".join(errors),
            })

        inner_dir = f"fitbase_client_photos_20260923_delta_75"
        manifest, archive_validation = write_zip(
            photo=photo, output=output, inner_dir=inner_dir,
            cutoff_at=ARGS.cutoff_at, targets=targets, photos=photos, filenames=names,
            missing_rows=missing_rows,
        )
        missing_fields = ["group", "client_id", "client_fio", "phone", "reason", "source_photo_errors"]
        atomic_write(missing_path, csv_bytes(missing_rows, missing_fields), overwrite=ARGS.overwrite)

        exported_counts = Counter(row["group"] for row in manifest)
        target_group_counts = Counter(target.group for target in targets.values())
        missing_counts = Counter(row["group"] for row in missing_rows)
        missing_reason_counts = Counter(row["reason"] for row in missing_rows)
        source_counts = Counter(row["photo_source"] for row in manifest)
        report = {
            "backup_set_uuid": ARGS.backup_set_uuid,
            "backup_finish_at": backup_finish_at,
            "cutoff_at": ARGS.cutoff_at,
            "cutoff_date": ARGS.cutoff_at[:10],
            "date_stamp": ARGS.cutoff_at[:10].replace("-", ""),
            "database": ARGS.database,
            "container": ARGS.container,
            "restorehistory": provenance,
            "target_source_counts": target_counts,
            "membership_workbook_checks": membership_checks,
            "target_counts_by_group": dict(sorted(target_group_counts.items())),
            "valid_phone_targets": len(valid_targets),
            "valid_phone_targets_by_group": dict(sorted(Counter(
                targets[client_id].group for client_id in valid_targets
            ).items())),
            "photos_exported": len(manifest),
            "photos_exported_by_group": dict(sorted(exported_counts.items())),
            "photos_missing": len(missing_rows),
            "photos_missing_by_group": dict(sorted(missing_counts.items())),
            "missing_reason_counts": dict(sorted(missing_reason_counts.items())),
            "excluded_ids_and_reasons": missing_rows,
            "shared_primary_phones": dict(sorted(shared_phones.items())),
            "source_candidate_rows": len(candidate_rows),
            "source_candidate_ids": len(candidates),
            "source_candidate_errors": candidate_errors,
            "selected_photo_sources": dict(sorted(source_counts.items())),
            "archive_manifest_rows": len(manifest),
            "archive_validation": archive_validation,
            "outputs": {
                "archive": str(output),
                "missing_photos_csv": str(missing_path),
                "report_json": str(report_json_path),
                "report_markdown": str(report_md_path),
                "log": str(log_path),
            },
            "input_sha256": {
                str(path.relative_to(ROOT)): sha256_file(path)
                for path in (LEADS_MANIFEST, MEMBERSHIP_MANIFEST, MEMBERSHIP_XLSX, Path(__file__).resolve())
            },
        }
        report_json = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
        markdown = textwrap.dedent(f"""\
            # Фотографии клиентов дельты за 23.09.2026

            Backup: `{ARGS.backup_set_uuid}`, завершён `{backup_finish_at}`.
            Единый срез `cutoff_at`: `{ARGS.cutoff_at}`.

            | Группа | Целей | Фото | Без фото |
            |---|---:|---:|---:|
            | `{PHOTO_GROUP}` | {target_group_counts.get(PHOTO_GROUP, 0)} | {exported_counts.get(PHOTO_GROUP, 0)} | {missing_counts.get(PHOTO_GROUP, 0)} |
            | `{REFUSER_GROUP}` | {target_group_counts.get(REFUSER_GROUP, 0)} | {exported_counts.get(REFUSER_GROUP, 0)} | {missing_counts.get(REFUSER_GROUP, 0)} |
            | **Всего** | **{len(targets)}** | **{len(manifest)}** | **{len(missing_rows)}** |

            Группа отказников основана на теге `{REFUSER_TAG}` и строках новых договоров.
            Для неё воронка не присваивалась, поскольку прежняя воронка неизвестна.

            Проверки: 75 целевых ID без пересечений; все 75 присутствуют в финальном XLSX абонементов;
            restorehistory UUID и время окончания совпали; CRC ZIP, уникальность путей,
            декодирование JPEG и сверка ID/хэшей manifest прошли.

            Архив: `{output}`
            Отсутствующие/непригодные фото и причины: `{missing_path}`
            JSON-отчёт: `{report_json_path}`
            Журнал: `{log_path}`
        """)
        atomic_write(report_json_path, report_json.encode("utf-8"), overwrite=ARGS.overwrite)
        atomic_write(report_md_path, markdown.encode("utf-8"), overwrite=ARGS.overwrite)
        logger.info("Published archive: %s", output)
        logger.info("Photos exported by group: %s", dict(sorted(exported_counts.items())))
        logger.info("Missing/excluded IDs by reason: %s", dict(sorted(missing_reason_counts.items())))
        logger.info("ZIP checks: %s", archive_validation)
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0
    finally:
        for handler in list(logger.handlers):
            handler.flush()
            handler.close()
            logger.removeHandler(handler)


if __name__ == "__main__":
    raise SystemExit(main())
