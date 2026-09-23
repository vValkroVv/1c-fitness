#!/usr/bin/env python3
"""Read-only evidence for the seven September membership-template conflicts.

Read the accepted workbook without saving it, reproduce candidate attributes
from already exported facts, and compare the underlying June/September rows.
Only SELECT queries are sent to SQL. No migration/staging writes are performed.
"""
from __future__ import annotations

import hashlib
import importlib.util
import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import sys
import time

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "end-to-end-xlsx"
OUT = ROOT / "output/20260921_template_audit"
ACCEPTED = ROOT / "output/20260630_delivery_funnel_labels_20260820/fitbase_import_shablony_abonementov_20260630.xlsx"
WORK = PACKAGE / "work/20260921_final"
JUNE = "FitnessRestored_20260630_original"
SEPTEMBER = "FitnessRestored_20260630_macos"
OLD_NAME = "Абонемент МУЛЬТИКАРТА 12 месяцев + подарок (Спецпредложение) рассрочка"
NEW_NAME = "Абонемент МУЛЬТИКАРТА 12 месяцев (Спецпредложение) рассрочка"
sys.path[:0] = [str(PACKAGE / "scripts"), str(ROOT / "scripts")]
from database import ConnectionSettings, DatabaseClient
from prepare_backup import query_dicts
from run_fitbase_migration import read_sql_password
from openpyxl import load_workbook


def load_builder():
    spec = importlib.util.spec_from_file_location("membership_template_audit_builder", PACKAGE / "scripts/19_build_membership_import_xlsx.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def save(name, data):
    (OUT / name).write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")


def quote(text):
    return "N'" + text.replace("'", "''") + "'"


def raw_contract_query(database, names):
    return f"""SELECT
 CONVERT(varchar(32),d._IDRRef,2) AS contract_ref,
 LTRIM(RTRIM(d._Number)) AS contract_id,
 CONVERT(varchar(32),p._IDRRef,2) AS product_ref,
 LTRIM(RTRIM(p._Code)) AS product_code,p._Description AS product_name,
 d._Date_Time AS document_date_raw,
 CASE WHEN d._Date_Time>'30000101' THEN DATEADD(year,-2000,d._Date_Time) ELSE d._Date_Time END AS document_date,
 CONVERT(int,d._Posted) AS posted,CONVERT(int,d._Marked) AS marked,
 d._Fld1481 AS document_duration,
 r._Fld3062 AS register_3062_raw,r._Fld3063 AS register_start_raw,r._Fld3064 AS register_end_raw,
 CASE WHEN r._Fld3063>'30000101' THEN DATEADD(year,-2000,r._Fld3063) ELSE r._Fld3063 END AS register_start,
 CASE WHEN r._Fld3064>'30000101' THEN DATEADD(year,-2000,r._Fld3064) ELSE r._Fld3064 END AS register_end,
 DATEDIFF(day,r._Fld3063,r._Fld3064)+1 AS interval_days,
 r._Fld3065 AS register_duration_days,r._Fld3068 AS register_freeze_days,
 r._Fld3069 AS register_guests,r._Fld3070 AS register_price,r._Fld3072 AS register_debt_candidate,
 st._Description AS status
FROM [{database}].dbo._Document163 d
JOIN [{database}].dbo._Reference72 p ON p._IDRRef=d._Fld1446RRef
LEFT JOIN [{database}].dbo._InfoRg3060 r ON r._Fld3061RRef=d._IDRRef
LEFT JOIN [{database}].dbo._Reference5062 st ON st._IDRRef=r._Fld5960RRef
WHERE LOWER(p._Description) IN ({','.join(quote(x.lower()) for x in names)})
ORDER BY d._Number OPTION (MAXDOP 2);\n"""


def source_attributes(row):
    """Reconstruct the relevant fields directly from the saved source record."""
    doc_duration = Decimal(row["document_duration"])
    days = row["interval_days"]
    if 0 < days < 28:
        duration = 0
    elif 0 < doc_duration <= 60:
        duration = int(doc_duration.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    else:
        duration = int((Decimal(days) / Decimal("30.4375")).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    return {
        "price": Decimal(row["register_price"]), "duration": duration,
        "visits": 12 if "сайкл 12 пос" in row["product_name"].lower() else None,
        "freeze": int(row["register_freeze_days"]) or None,
        "branches_access": "Все" if "мультикарта" in row["product_name"].lower() else "Продажа",
    }


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    builder = load_builder()
    conflicts = json.loads((ROOT / "logs/20260921_template_source_conflicts.json").read_text())
    targets = {item["name"] for item in conflicts}
    workbook = load_workbook(ACCEPTED, read_only=True, data_only=True)
    accepted = []
    for sheet in workbook:
        headers = None
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if "name" in row and "price" in row:
                headers = list(row)
                continue
            if not headers:
                continue
            values = dict(zip(headers, row))
            name = str(values.get("name") or "")
            if builder.normalize_key(name) in targets or name == OLD_NAME:
                accepted.append({"sheet": sheet.title, "range": f"A{row_number}:L{row_number}", "values": values})
    workbook.close()
    assert len(accepted) == 7, f"Expected seven accepted rows, got {len(accepted)}"
    save("accepted_rows.json", {"path": str(ACCEPTED.relative_to(ROOT)), "sha256": hashlib.sha256(ACCEPTED.read_bytes()).hexdigest(), "rows": accepted})

    source_clients = builder.read_source_clients(WORK / "owner/fitbase_active_clients_import_zayavki_20260921_all_funnels.xlsx")
    source_clients.update(builder.read_refuser_clients(WORK / "owner/csv/new_application_refusers.csv"))
    facts = builder.read_facts(WORK / "imports/staging/membership_import_facts.tsv")
    excluded, _ = builder.find_contact_next_exclusions(facts)
    candidates = []
    for fact in facts:
        name = fact["subscription_name"].strip()
        if builder.normalize_key(name) not in targets:
            continue
        if fact["client_id"].strip() not in source_clients or fact["subscription_ref"].strip() in excluded:
            continue
        duration, duration_source = builder.compute_duration_months(fact)
        price, price_source = builder.choose_financial_price(fact)
        override = builder.business_zero_override_reason(fact, price)
        if override:
            price, price_source = 0, override
        candidates.append({
            "name": name,"contract_id": fact["document_number"].strip(),
            "contract_ref": fact["subscription_ref"],"product_ref": fact["product_ref"],
            "price": builder.excel_number(price),"price_source": price_source,
            "duration": duration,"duration_source": duration_source,
            "visits": builder.parse_template_visits(name) if builder.is_visit_limited_membership(fact,name) else None,
            "freeze": builder.int_or_blank(builder.decimal_value(fact["rg_freeze_days"])),
            "branches_access": "Все" if "мультикарта" in builder.normalize_key(name) else "Продажа",
        })
    save("current_candidates.json", candidates)

    names = [row["values"]["name"] for row in accepted] + [NEW_NAME]
    password = read_sql_password(ROOT / "tmp/macos-backup/mssql-fitness-macos.env")
    raw = {}
    with DatabaseClient(ConnectionSettings(server="127.0.0.1", port=11435, database=SEPTEMBER, user="sa", password=password, query_timeout_seconds=180)) as db:
        for key, database in [("june", JUNE), ("september", SEPTEMBER)]:
            sql = raw_contract_query(database, names)
            (OUT / f"{key}_contracts.sql").write_text(sql, encoding="utf-8")
            raw[key] = query_dicts(db, sql)
            save(f"{key}_contracts.json", raw[key])
            print(f"{key}: {len(raw[key])} source contracts", flush=True)

    june_index = {r["contract_ref"]: r for r in raw["june"]}
    september_index = {r["contract_ref"]: r for r in raw["september"]}
    assert len(june_index) == len(raw["june"]), "June register returned duplicate contract rows"
    assert len(september_index) == len(raw["september"]), "September register returned duplicate contract rows"
    fields = ["price", "duration", "visits", "freeze", "branches_access"]
    resolutions = []
    for accepted_row in accepted:
        values = accepted_row["values"]
        name = NEW_NAME if values["name"] == OLD_NAME else values["name"]
        conflict = next(item for item in conflicts if item["name"] == builder.normalize_key(name))
        matching = [row for row in candidates if builder.normalize_key(row["name"]) == builder.normalize_key(name) and all(row[field] == values[field] for field in fields)]
        matching.sort(key=lambda row: (row["contract_ref"] not in june_index, row["contract_id"]))
        existing_id = conflict["source_contract_id"]
        selected = next((row for row in candidates if row["contract_id"] == existing_id), None) if existing_id else (matching[0] if matching else None)
        assert selected, f"Missing source contract for {name}"
        ref = selected["contract_ref"]
        prior = june_index.get(ref)
        current = september_index[ref]
        assert prior is not None, f"Selected contract missing in June: {selected['contract_id']}"
        assert current["posted"] == 1 and current["marked"] == 0, f"Source contract is not posted and active: {name}"
        assert source_attributes(prior) == {field: values[field] for field in fields}, f"June source does not support accepted attributes: {name}"
        assert source_attributes(current) == {field: selected[field] for field in fields}, f"Current candidate differs from raw source: {name}"
        changes = {key: {"june": prior[key], "september": current[key]} for key in current if prior and prior[key] != current[key]}
        resolutions.append({
            "current_name": name,"accepted_name": values["name"],
            "accepted_workbook_range": accepted_row["range"],
            "accepted_attributes": {field: values[field] for field in fields},
            "source_contract_id": selected["contract_id"],"source_contract_ref": ref,
            "selection": "preserve_existing_historical_anchor" if existing_id else "exact_variant_existing_in_june_preferred",
            "matching_current_contracts": [row["contract_id"] for row in matching],
            "source_present_in_june": prior is not None,"current_candidate": selected,
            "june_source": prior,"september_source": current,"source_changes": changes,
        })
    save("resolutions.json", resolutions)
    status = {"status": "PASS", "accepted_rows": len(accepted), "resolutions": len(resolutions), "new_anchor_count": sum(r["selection"].startswith("exact") for r in resolutions), "historical_anchor_count": sum(r["selection"].startswith("preserve") for r in resolutions), "cutoff_at": "2026-09-21 20:12:12", "sources_read_only": True}
    save("summary.json", status)
    print(json.dumps(status, ensure_ascii=False), flush=True)
    for result in resolutions:
        print(json.dumps({"name": result["current_name"], "source_contract_id": result["source_contract_id"], "june": result["source_present_in_june"], "changes": result["source_changes"]}, ensure_ascii=False, default=str), flush=True)
    return 0


if __name__ == "__main__":
    started = time.monotonic()
    result = main()
    print(f"Completed in {time.monotonic() - started:.1f}s", flush=True)
    raise SystemExit(result)
