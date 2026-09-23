#!/usr/bin/env python3
"""Recover the exact 371 phone-deduplicated clients from the September 21 run.

Reuse the accepted builders for business values, saved facts for the cutoff,
and Artifact Tool for XLSX authoring. SQL is read only, for original photo BLOBs.
"""

import argparse
import csv
import hashlib
import importlib.util
import io
import json
import re
import subprocess
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PIPELINE = ROOT / "end-to-end-xlsx"
SOURCE = PIPELINE / "work/20260921_final_v2"
ORIGINAL = ROOT / "output/20260921_fitbase_for_customer"
WORK = PIPELINE / "work/20260921_phone_dedup_supplement_371"
OUTPUT = ROOT / "output/20260921_phone_dedup_supplement_371"
CUTOFF = "2026-09-21 20:12:12"
BACKUP = "2026-09-20 20:12:12"
DATABASE = "FitnessRestored_20260630_macos"
CONTAINER = "mssql-fitness-2022"
CLIENT_NAME = "fitbase_active_clients_import_zayavki_20260921_all_funnels.xlsx"


def module(filename):
    spec = importlib.util.spec_from_file_location(filename.removesuffix(".py"), PIPELINE / "scripts" / filename)
    result = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = result
    spec.loader.exec_module(result)
    return result


def digest(path):
    with path.open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def json_write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str) + "\n")


def csv_read(path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def csv_write(path, rows, fields):
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def workbook_rows(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        stream = sheet.iter_rows(values_only=True)
        headers, russian = list(next(stream)), list(next(stream))
        rows = [dict(zip(headers, values)) for values in stream if any(v is not None for v in values)]
        return headers, russian, rows
    finally:
        book.close()


def selected_clients():
    excluded = csv_read(SOURCE / "owner/reports/phone_deduplication_removed_clients.csv")
    selected = [r for r in excluded if r["funnel"] in {"Действующие клиенты", "Реактивация"}]
    assert Counter(r["funnel"] for r in selected) == {"Действующие клиенты": 68, "Реактивация": 303}
    assert len({r["client_id"] for r in selected}) == 371
    return selected


def prepare():
    WORK.mkdir(parents=True, exist_ok=True)
    selected = selected_clients()
    ids = {r["client_id"] for r in selected}
    original_headers, original_russian, original_clients = workbook_rows(ORIGINAL / CLIENT_NAME)
    assert not ids & {r["client_id"] for r in original_clients}
    main = module("12_build_part2_three_funnel_xlsx.py")
    combined = module("17_build_part2_combined_xlsx.py")
    membership = module("19_build_membership_import_xlsx.py")
    services = module("23_build_services_import_xlsx.py")
    stage = csv_read(SOURCE / "owner/staging/final_funnel_clients.csv")
    rows = [r for r in stage if r["client_id"] in ids]
    assert len(rows) == 371 and {r["client_id"] for r in rows} == ids
    excluded_map = {r["client_id"]: r for r in selected}
    for row in rows:
        old = excluded_map[row["client_id"]]
        assert row["cutoff_date"] == CUTOFF[:10]
        for field in ("client_ref", "client_fio", "phones", "funnel", "selected_subscription_ref"):
            assert row[field] == old[field], (row["client_id"], field)
    meta = csv_read(SOURCE / "owner/staging/staging_run_metadata.csv")
    assert len(meta) == 1 and meta[0]["cutoff_at"][:19] == CUTOFF
    assert meta[0]["backup_finish_at"][:19] == BACKUP
    main.assign_branches(rows, main.load_branches(PIPELINE / "config/branches_by_club.yml"))
    rows = main.sort_rows(combined.apply_fitbase_labels(rows, combined.CUSTOMER_SINGLE_STAGE_MODE))
    clients = [{
        "client_id": r["client_id"], "phone": r["phones"], "client_fio": r["client_fio"],
        "email": r["email"], "funnel": r["funnel"], "funnel_step": r["funnel_step"],
        "budget": 0, "create_date": main.parse_date(r["create_date"]),
        "manager": r["manager"], "филиал": r["branch"],
    } for r in rows]
    source_clients = {r["client_id"]: membership.SourceClient(
        r["client_id"], r["phone"], r["client_fio"], r["create_date"], r["manager"], r["филиал"]
    ) for r in clients}
    facts = membership.read_facts(SOURCE / "imports/staging/membership_import_facts.tsv")
    assert {f["cutoff_at"] for f in facts} == {CUTOFF}
    template_name = "fitbase_import_shablony_abonementov_20260921.xlsx"
    _, _, original_templates = workbook_rows(ORIGINAL / template_name)
    # The accepted full delivery is authoritative for template parameters.
    # Cohort selection must not choose a different price/duration variant.
    decisions = membership.read_template_canonicalizations(PIPELINE / "config/membership_template_canonicalization.csv")
    for row in original_templates:
        decisions[membership.normalize_key(row["name"])] = membership.TemplateCanonicalization(
            row["name"], row["branches_access"], row["price"], row["duration"], row["visits"],
            row["freeze"], "", "Accepted September 21 delivery", "accepted",
            "Keep the full-delivery template unchanged in the supplement", True,
        )
    memberships, derived_templates, uncertainties, excluded_facts, counters = membership.build_rows(
        source_clients, membership.read_cards(SOURCE / "owner/staging"), facts, decisions)
    assert all(r["client_id"] in ids for r in memberships)
    assert len({r["contract_id"] for r in memberships}) == len(memberships)
    assert "00000151350" not in {r["contract_id"] for r in memberships}
    active_ids = {r["client_id"] for r in clients if r["funnel"] == "Действующие абонементы"}
    assert active_ids <= {r["client_id"] for r in memberships}
    names = {r["contract_name"] for r in memberships}
    template_map = {r["name"]: r for r in original_templates}
    for row in derived_templates:
        template_map.setdefault(row["name"], row)
    membership_templates = [template_map[name] for name in sorted(names)]
    service_facts = services.read_facts(SOURCE / "imports/staging/services_import_facts.tsv")
    assert {f["cutoff_at"] for f in service_facts} == {CUTOFF}
    # Select services in the full delivery population plus the recovered clients,
    # then restrict output to the requested IDs. This preserves global fallback rules.
    service_sources = services.read_source_clients(ORIGINAL / CLIENT_NAME)
    service_sources.update({r["client_id"]: services.SourceClient(
        r["client_id"], r["phone"], r["client_fio"], r["create_date"], r["manager"],
        r["филиал"], r["funnel"], r["funnel_step"],
    ) for r in clients})
    service_rows, _, service_issues, coverage, _ = services.build_rows(
        service_sources, services.read_service_names(PIPELINE / "templates/services_required.xlsx"),
        service_facts, manager_pools=main.load_managers(PIPELINE / "config/managers_by_club.yml"))
    service_rows = [r for r in service_rows if r["client_id"] in ids]
    service_template_name = "fitbase_import_shablony_uslug_20260921.xlsx"
    _, _, service_templates = workbook_rows(ORIGINAL / service_template_name)
    service_names = {r["service_name"] for r in service_rows}
    service_templates = [r for r in service_templates if r["name"] in service_names]
    assert {r["name"] for r in service_templates} == service_names
    packages = [
        (CLIENT_NAME, original_headers, original_russian, clients),
        ("fitbase_import_abonementy_clientov_20260921.xlsx", membership.CLIENT_HEADERS, membership.CLIENT_RUS_HEADERS, memberships),
        (template_name, membership.TEMPLATE_HEADERS, membership.TEMPLATE_RUS_HEADERS, membership_templates),
        ("fitbase_import_uslugi_clientov_20260921.xlsx", services.CLIENT_HEADERS, services.CLIENT_RUS_HEADERS, service_rows),
        (service_template_name, services.TEMPLATE_HEADERS, services.TEMPLATE_RUS_HEADERS, service_templates),
    ]
    specs = []
    for name, headers, russian, data in packages:
        specs.append({"name": name, "headers": headers, "russian_headers": russian,
                      "rows": [[r.get(h) for h in headers] for r in data]})
    summary = {
        "cutoff_at": CUTOFF, "backup_finish_at": BACKUP, "clients": len(clients),
        "funnels": dict(Counter(r["funnel"] for r in clients)),
        "membership_rows": len(memberships), "membership_clients": len({r["client_id"] for r in memberships}),
        "clients_without_importable_memberships": sorted(ids - {r["client_id"] for r in memberships}),
        "membership_templates": len(membership_templates), "services": len(service_rows),
        "service_clients": len({r["client_id"] for r in service_rows}), "service_templates": len(service_templates),
        "service_row_kinds": dict(Counter(r["_row_kind"] for r in service_rows)),
        "problem4_omitted": "Contract 151350 belongs to client 000036376, outside the requested 371 IDs",
        "membership_counters": counters,
    }
    json_write(WORK / "workbooks.json", specs)
    json_write(WORK / "summary.json", summary)
    json_write(WORK / "membership_uncertainties.json", uncertainties)
    json_write(WORK / "membership_exclusions.json", [r for r in excluded_facts if r.get("client_id") in ids])
    json_write(WORK / "service_coverage.json", coverage)
    csv_write(WORK / "selected_clients.csv", selected, list(selected[0]))
    input_files = [SOURCE / "owner/reports/phone_deduplication_removed_clients.csv",
                   SOURCE / "owner/staging/final_funnel_clients.csv",
                   SOURCE / "owner/staging/staging_run_metadata.csv",
                   SOURCE / "imports/staging/membership_import_facts.tsv",
                   SOURCE / "imports/staging/services_import_facts.tsv"]
    input_files += sorted(ORIGINAL.iterdir())
    input_files += [PIPELINE / "scripts" / name for name in (
        "12_build_part2_three_funnel_xlsx.py", "17_build_part2_combined_xlsx.py",
        "19_build_membership_import_xlsx.py", "23_build_services_import_xlsx.py",
        "41_export_active_client_photos_zip.py")]
    json_write(WORK / "source_hashes.json", {str(p.relative_to(ROOT)): digest(p) for p in input_files})
    print(json.dumps(summary, ensure_ascii=False, default=str))


def photo_query(photo, ids, include_blob):
    query = photo.sql_selected_union(DATABASE, include_blob=include_blob, cutoff_date="20260921")
    join = f"JOIN [{DATABASE}].fitbase_part2.final_funnel_clients AS a ON a.client_ref = CONVERT(varchar(32), c._IDRRef, 2) "
    condition = "a.cutoff_date = CONVERT(date, '20260921', 112)"
    assert query.count(join) == query.count(condition) == 2
    assert all(re.fullmatch(r"[0-9]{9}", client_id) for client_id in ids)
    # Read the immutable restored source by the historical CSV IDs, without
    # consulting or changing the current September 22 staging tables.
    query = query.replace(join, "").replace(condition, "c._Code IN (" + ",".join("'" + i + "'" for i in sorted(ids)) + ")")
    return query


def photos():
    photo = module("41_export_active_client_photos_zip.py")
    selected = selected_clients()
    ids = {r["client_id"] for r in selected}
    # Verify that photo BLOBs still come from the same restored September backup.
    query = f"""SET NOCOUNT ON; SELECT TOP (1) CONVERT(varchar(36), bs.backup_set_uuid) AS backup_set_uuid,
    CONVERT(varchar(19), bs.backup_finish_date,120) AS backup_finish_at
    FROM msdb.dbo.restorehistory rh JOIN msdb.dbo.backupset bs ON bs.backup_set_id=rh.backup_set_id
    WHERE rh.destination_database_name=N'{DATABASE}' AND rh.restore_type='D'
    ORDER BY rh.restore_date DESC FOR JSON PATH;"""
    process = subprocess.run(["docker", "exec", CONTAINER, "/bin/bash", "-lc",
        'exec /opt/mssql-tools18/bin/sqlcmd -S localhost -U sa -P "$MSSQL_SA_PASSWORD" -C "$@"',
        "sqlcmd", "-d", DATABASE, "-b", "-Q", query, "-y", "0", "-w", "65535"],
        text=True, capture_output=True, check=True)
    text = "".join(process.stdout.splitlines())
    provenance = json.loads(text[text.index("[{"):text.rindex("}]")+2])
    assert provenance[0]["backup_set_uuid"].lower() == "8e76ec9f-004f-4291-9842-92017c8e9f53"
    assert provenance[0]["backup_finish_at"] == BACKUP
    clients = photo.load_delivered_clients(OUTPUT / CLIENT_NAME)
    assert set(clients) == ids
    metadata = photo.load_metadata(CONTAINER, photo_query(photo, ids, False))
    eligible, exclusions = photo.select_eligible_photos(metadata, clients)
    owners = defaultdict(set)
    original = photo.load_delivered_clients(ORIGINAL / CLIENT_NAME)
    for client in [*original.values(), *clients.values()]:
        for phone_number in client.normalized_phones:
            owners[phone_number].add(client.client_id)
    assignments = {}
    for client_id, item in eligible.items():
        unique = [number for number in item.normalized_phones if owners[number] == {client_id}]
        number = unique[0] if unique else item.normalized_phones[0]
        status = ("exact_primary_phone" if number == item.normalized_phones[0] else "exact_alternate_phone") if unique else "duplicate_phone_client_id_suffix"
        basename = number if unique else f"{number}__{client_id}"
        assignments[client_id] = photo.NameAssignment(basename, number, status)
    result = photo.write_archive(
        container=CONTAINER, query=photo_query(photo, ids, True), metadata=eligible,
        assignments=assignments, clients=clients, exclusions=exclusions,
        output_path=OUTPUT / "fitbase_client_photos_20260921.zip",
        inner_dir="fitbase_client_photos_20260921_supplement_371",
        cutoff_at=CUTOFF, backup_finish_at=BACKUP, expected_count=len(eligible), overwrite=False)
    result["source_backup"] = provenance[0]
    json_write(WORK / "photos.json", result)
    print(json.dumps(result, ensure_ascii=False))


def normalized(value):
    if value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def verify():
    ids = {r["client_id"] for r in selected_clients()}
    specs = json.loads((WORK / "workbooks.json").read_text())
    summary = json.loads((WORK / "summary.json").read_text())
    photo_report = json.loads((WORK / "photos.json").read_text())
    files = []
    client_map = {}
    for spec in specs:
        path = OUTPUT / spec["name"]
        headers, russian, rows = workbook_rows(path)
        assert headers == spec["headers"] and russian == spec["russian_headers"], path.name
        actual = [[normalized(r[h]) for h in headers] for r in rows]
        expected = [[normalized(v) for v in row] for row in spec["rows"]]
        assert actual == expected, path.name
        if "client_id" in headers:
            assert {r["client_id"] for r in rows} <= ids
            assert all(isinstance(r["client_id"], str) and re.fullmatch(r"\d{9}", r["client_id"]) for r in rows)
        if spec["name"] == CLIENT_NAME:
            assert {r["client_id"] for r in rows} == ids and len(rows) == 371
            assert Counter(r["funnel"] for r in rows) == {"Действующие абонементы": 68, "Реактивация": 303}
            client_map = {r["client_id"]: r for r in rows}
        elif "client_id" in headers:
            for row in rows:
                for field in ("phone", "client_fio", "manager"):
                    assert normalized(row[field]) == normalized(client_map[row["client_id"]][field])
        files.append({"name": path.name, "rows": len(rows), "bytes": path.stat().st_size, "sha256": digest(path)})
    zip_path = OUTPUT / "fitbase_client_photos_20260921.zip"
    with zipfile.ZipFile(zip_path) as archive:
        manifest_name = next(n for n in archive.namelist() if n.endswith("/_reports/manifest.csv"))
        manifest = list(csv.DictReader(io.StringIO(archive.read(manifest_name).decode("utf-8-sig"))))
        excluded_name = manifest_name.replace("manifest.csv", "excluded_clients.csv")
        exclusions = list(csv.DictReader(io.StringIO(archive.read(excluded_name).decode("utf-8-sig"))))
        assert {r["client_id"] for r in manifest}.isdisjoint({r["client_id"] for r in exclusions})
        assert {r["client_id"] for r in manifest + exclusions} == ids
        with zipfile.ZipFile(ORIGINAL / zip_path.name) as old:
            original_names = {Path(n).name for n in old.namelist() if n.endswith(".jpg")}
        assert not {r["filename"] for r in manifest} & original_names
    files.append({"name": zip_path.name, "photos": len(manifest), "bytes": zip_path.stat().st_size, "sha256": digest(zip_path)})
    for relative, checksum in json.loads((WORK / "source_hashes.json").read_text()).items():
        assert digest(ROOT / relative) == checksum, f"Source changed: {relative}"
    assert len(list(OUTPUT.glob("*.xlsx"))) == 5
    assert not list(OUTPUT.glob("*plastic*")) and not list(OUTPUT.glob("problem*"))
    result = {"status": "PASS", "summary": summary, "photos": photo_report, "files": files,
              "checks": ["exact 371 source IDs", "68 active / 303 reactivation", "all saved cells match planned values",
                         "client IDs remain text", "phones/names/managers consistent", "no unrelated client IDs",
                         "common September 21 cutoff", "original sources unchanged", "photo ID coverage and no old filename collisions"]}
    json_write(WORK / "validation.json", result)
    print(json.dumps({"status": "PASS", "files": files}, ensure_ascii=False))


def finalize():
    verify()
    validation = json.loads((WORK / "validation.json").read_text())
    assert validation["status"] == "PASS"
    text = """Дополнительная выгрузка Fitbase за 21 сентября 2026 года

Ровно 371 клиент из отчёта об исключении по общим телефонам:
68 — «Действующие абонементы», 303 — «Реактивация».
Срез всех данных: 21.09.2026 20:12:12.
Фактическое окончание backup: 20.09.2026 20:12:12.

Состав папки
1. Заявки — 371 клиент.
2. Абонементы клиентов — 879 записей, 367 клиентов.
3. Шаблоны абонементов — 59 используемых шаблонов.
4. Услуги клиентов — 10 записей, 7 клиентов.
5. Шаблоны услуг — 5 используемых шаблонов.
6. Архив фотографий — 244 JPEG.

Файл пластиковых карт исключён по запросу. Поле «Номер карты» в абонементах
сохранено, как в основной выгрузке. Отдельный problem-файл основной поставки
к этим 371 клиенту не относится и в папку не включён.

Особенности импорта
Повторное удаление по совпадающим телефонам не выполнялось. Все исходные
номера телефонов и ID этих 371 клиентов сохранены. Совпадения телефонов
с уже переданными клиентами остаются; сопоставление карточек нужно выполнять
по client_id. Данные других клиентов в импортные файлы не добавлены.

У четырёх клиентов реактивации исторические договоры не проходят правила
отбора абонементов основной поставки. Все четверо присутствуют в заявках:
000027427 — Свек Де Педраза еунисе ноеми;
000036594 — Цыбульник Татьяна Степановна;
000065047 — Елизавета;
000065636 — Казакова Валентина Витальевна.
У всех 68 действующих клиентов абонементы включены.

Фотографии
Для 244 клиентов найден и проверен JPEG, для остальных 127 в backup
не найден подходящий фотофайл. Полный список исключений есть внутри ZIP.
84 фотографии имеют имя по уникальному телефону; в 34 из этих случаев
использован альтернативный номер самого клиента.
У 160 клиентов все подходящие телефоны общие с другими карточками.
Имена их фото: <телефон>__<client_id>.jpg. Сопоставлять такие фото только
по телефону неоднозначно: используйте client_id из _reports/manifest.csv.
Список этих 160 файлов — _reports/phone_filename_exceptions.csv.
Имена JPEG не пересекаются с фотоархивом основной поставки 21 сентября.

Проверено: точное совпадение списка 371 ID с отчётом об исключениях,
распределение 68/303, все значения сохранённых XLSX, текстовый тип ID,
единый срез, соответствие телефонов и менеджеров, CRC и SHA-256 всех JPEG.
Это подготовленные файлы для дополнительного импорта; загрузка в Fitbase
в рамках этой работы не выполнялась.
"""
    readme = OUTPUT / "README.txt"
    readme.write_text(text, encoding="utf-8-sig")
    expected_names = {r["name"] for r in validation["files"]} | {readme.name}
    assert {p.name for p in OUTPUT.iterdir()} == expected_names
    validation["files"].append({"name": readme.name, "bytes": readme.stat().st_size, "sha256": digest(readme)})
    json_write(WORK / "delivery_manifest.json", validation)
    (WORK / "READY.txt").write_text("PASS\n" + CUTOFF + "\n")
    print(str(OUTPUT))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("step", choices=["prepare", "photos", "verify", "finalize"])
    globals()[parser.parse_args().step]()
