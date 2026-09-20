#!/usr/bin/env python3
"""Export the delivered active/reactivation clients' valid-phone photos as JPEG.

The client workbook is the authority for IDs, funnel scope and phone numbers.
The restored database supplies the wrapped 1C image payloads. Missing/invalid
phones, absent photos and clients excluded from that workbook are reported and
never enter the ZIP. Shared valid phones retain explicit client-ID suffixes.

``cutoff_at`` is the effective business timestamp; ``backup_finish_at`` records
the actual backup timestamp independently. The pipeline verifies backup identity
before invoking this exporter and all SQL export layers must match cutoff_at.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import struct
import subprocess
import sys
import zlib
import zipfile
from collections import Counter
from dataclasses import dataclass, replace
from datetime import datetime
from pathlib import Path
from typing import Sequence

from openpyxl import load_workbook
from PIL import Image, ImageOps


ROOT = Path(__file__).resolve().parents[1]
ALLOWED_FUNNELS = frozenset({"Действующие клиенты", "Действующие абонементы", "Реактивация"})

PHONE_SPLIT_RE = re.compile(r"[,;]\s*")
SAFE_IDENTIFIER_RE = re.compile(r"^[A-Za-z0-9_]+$")
SAFE_CONTAINER_RE = re.compile(r"^[A-Za-z0-9_.-]+$")
IMAGE_SIGNATURES: tuple[tuple[bytes, str], ...] = (
    (bytes.fromhex("FFD8FF"), "jpg"),
    (bytes.fromhex("89504E470D0A1A0A"), "png"),
    (b"BM", "bmp"),
    (b"GIF8", "gif"),
)


@dataclass(frozen=True)
class PhotoMeta:
    client_id: str
    raw_phone: str
    source: str
    metadata_extension: str
    normalized_phones: tuple[str, ...]


@dataclass(frozen=True)
class DeliveredClient:
    client_id: str
    raw_phone: str
    funnel: str
    normalized_phones: tuple[str, ...]


@dataclass(frozen=True)
class NameAssignment:
    basename: str
    assigned_phone: str
    status: str


def require_safe_identifier(value: str, label: str) -> str:
    """Allow only the identifier subset needed by the known runtime names."""

    if not SAFE_IDENTIFIER_RE.fullmatch(value):
        raise ValueError(f"Unsafe {label}: {value!r}")
    return value


def require_safe_container(value: str) -> str:
    """Validate a Docker container name passed as a direct argv element."""

    if not SAFE_CONTAINER_RE.fullmatch(value):
        raise ValueError(f"Unsafe container name: {value!r}")
    return value


def normalize_phone_token(value: str) -> str:
    """Normalize Russian mobile/telephone numbers like the accepted XLSX pipeline."""

    digits = re.sub(r"\D+", "", value or "")
    if len(digits) == 11 and digits[0] in {"7", "8"}:
        return "7" + digits[1:]
    if len(digits) == 10:
        return "7" + digits
    return ""


def normalize_phones(value: str) -> tuple[str, ...]:
    """Return unique valid phones while preserving their order in the client card."""

    result: list[str] = []
    for token in PHONE_SPLIT_RE.split(value or ""):
        phone = normalize_phone_token(token)
        if phone and phone not in result:
            result.append(phone)
    return tuple(result)


def report_text(value: str) -> str:
    """Keep CSV rows one-line and make embedded NUL values visible."""

    return (value or "").replace("\x00", "\\0").replace("\t", " ").replace("\r", " ").replace("\n", " ")


def sql_selected_union(database: str, *, include_blob: bool, cutoff_date: str) -> str:
    """Build photo candidates for every internal funnel at the requested cutoff.

    The delivered workbook, including its reclassification and deduplication,
    determines the final scope in Python; internal SQL funnel labels do not.

    Newer records expose their main file through ``_Fld3834 -> Reference65``.
    Older records expose the in-backup migrated copy through
    ``_Fld9193 -> Reference115``.  Both payloads live in ``_InfoRg6767``.
    """

    db = f"[{require_safe_identifier(database, 'database name')}]"
    phone_sql = (
        "REPLACE(REPLACE(REPLACE(c._Fld3832, CHAR(9), N' '), "
        "CHAR(10), N' '), CHAR(13), N' ')"
    )
    blob_column = ", i._Fld6769 AS blob_value" if include_blob else ""

    newer = f"""
        SELECT
            c._Code AS client_id,
            {phone_sql} AS raw_phone,
            N'Reference65_main' AS photo_source,
            LOWER(f._Fld3884) AS metadata_extension
            {blob_column}
        FROM {db}.dbo._Reference64 AS c
        JOIN {db}.fitbase_part2.final_funnel_clients AS a
          ON a.client_ref = CONVERT(varchar(32), c._IDRRef, 2)
        JOIN {db}.dbo._Reference65 AS f
          ON c._Fld3834_RTRef = 0x00000041
         AND f._IDRRef = c._Fld3834_RRRef
        JOIN {db}.dbo._InfoRg6767 AS i
          ON i._Fld6768_RTRef = 0x00000041
         AND i._Fld6768_RRRef = f._IDRRef
        WHERE a.cutoff_date = CONVERT(date, '{cutoff_date}', 112)
          AND i._Fld6769 IS NOT NULL AND DATALENGTH(i._Fld6769) > 0
          AND LOWER(f._Fld3884) IN (N'jpg', N'jpeg', N'png', N'bmp', N'gif')
    """
    older = f"""
        SELECT
            c._Code AS client_id,
            {phone_sql} AS raw_phone,
            N'Reference115_copy' AS photo_source,
            LOWER(f._Fld4606) AS metadata_extension
            {blob_column}
        FROM {db}.dbo._Reference64 AS c
        JOIN {db}.fitbase_part2.final_funnel_clients AS a
          ON a.client_ref = CONVERT(varchar(32), c._IDRRef, 2)
        JOIN {db}.dbo._Reference115 AS f
          ON c._Fld3834_RTRef = 0x00000073
         AND c._Fld9193_RTRef = 0x00000073
         AND f._IDRRef = c._Fld9193_RRRef
        JOIN {db}.dbo._InfoRg6767 AS i
          ON i._Fld6768_RTRef = 0x00000073
         AND i._Fld6768_RRRef = f._IDRRef
        WHERE a.cutoff_date = CONVERT(date, '{cutoff_date}', 112)
          AND i._Fld6769 IS NOT NULL AND DATALENGTH(i._Fld6769) > 0
          AND LOWER(f._Fld4606) IN (N'jpg', N'jpeg', N'png', N'bmp', N'gif')
    """

    columns = "client_id, raw_phone, photo_source, metadata_extension"
    if include_blob:
        columns += ", blob_value"
    query = f"""
        SELECT {columns}
        FROM (
            {newer}
            UNION ALL
            {older}
        ) AS selected
        ORDER BY TRY_CONVERT(bigint, client_id), client_id
    """
    return " ".join(query.split())


def bcp_command(container: str, query: str) -> list[str]:
    """Run BCP inside the existing SQL Server container and stream UTF-8 rows."""

    require_safe_container(container)
    return [
        "docker",
        "exec",
        container,
        "/bin/bash",
        "-lc",
        'exec /opt/mssql-tools18/bin/bcp "$@" -S localhost -U sa -P "$MSSQL_SA_PASSWORD" -u',
        "bcp",
        query,
        "queryout",
        "/dev/stdout",
        "-c",
        "-C",
        "65001",
        "-t",
        "\t",
        "-r",
        "\n",
    ]


def start_bcp(container: str, query: str) -> subprocess.Popen[bytes]:
    """Start BCP with stdout available as a row stream."""

    process = subprocess.Popen(
        bcp_command(container, query),
        cwd=ROOT,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    if process.stdout is None or process.stderr is None:  # pragma: no cover - defensive
        process.kill()
        raise RuntimeError("Could not open BCP pipes")
    return process


def finish_bcp(process: subprocess.Popen[bytes], unexpected_lines: Sequence[bytes]) -> None:
    """Check BCP completion after its data stream has been consumed."""

    stderr = process.stderr.read().decode("utf-8", errors="replace")
    return_code = process.wait()
    if return_code != 0:
        details = stderr.strip() or b"\n".join(unexpected_lines[-10:]).decode("utf-8", errors="replace")
        raise RuntimeError(f"BCP failed with exit code {return_code}: {details}")
    if unexpected_lines:
        sample = b"\n".join(unexpected_lines[:3]).decode("utf-8", errors="replace")
        raise RuntimeError(f"Unexpected BCP output while reading data rows: {sample}")


def is_bcp_status_line(line: bytes) -> bool:
    stripped = line.strip()
    return (
        not stripped
        or stripped.startswith(b"Starting copy")
        or b" rows copied." in stripped
        or b" rows successfully bulk-copied to host-file." in stripped
        or stripped.startswith(b"Network packet size")
        or stripped.startswith(b"Clock Time")
    )


def close_bcp(process: subprocess.Popen[bytes]) -> None:
    """Always reap the SQL child, including malformed-row and decoding failures."""

    if process.poll() is None:
        process.kill()
    process.wait()
    if process.stdout is not None:
        process.stdout.close()
    if process.stderr is not None:
        process.stderr.close()


def load_metadata(container: str, query: str) -> dict[str, PhotoMeta]:
    """Read the small metadata pass used to assign unique archive filenames."""

    process = start_bcp(container, query)
    metadata: dict[str, PhotoMeta] = {}
    unexpected: list[bytes] = []
    try:
        assert process.stdout is not None
        for raw_line in process.stdout:
            line = raw_line.rstrip(b"\r\n")
            fields = line.split(b"\t")
            if len(fields) != 4:
                if not is_bcp_status_line(line):
                    unexpected.append(line[:500])
                continue
            client_id, raw_phone, source, extension = (field.decode("utf-8", errors="strict") for field in fields)
            if client_id in metadata:
                raise RuntimeError(f"Client {client_id} was selected more than once")
            metadata[client_id] = PhotoMeta(
                client_id=client_id, raw_phone=raw_phone, source=source,
                metadata_extension=extension, normalized_phones=normalize_phones(raw_phone),
            )
        finish_bcp(process, unexpected)
        return metadata
    finally:
        close_bcp(process)


def load_delivered_clients(path: Path) -> dict[str, DeliveredClient]:
    """Read actual delivery IDs and phones, preserving leading zeros in IDs."""

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, ())
        indexes = {str(value).strip(): index for index, value in enumerate(header)}
        required = {"client_id", "phone", "funnel"}
        if not required.issubset(indexes):
            raise ValueError(f"Client workbook must contain columns {sorted(required)}: {path}")
        russian_header = next(rows, ())
        if not russian_header or str(russian_header[indexes["client_id"]]).strip() != "Внутренний номер клиента":
            raise ValueError(f"Client workbook is missing its expected second header row: {path}")
        clients: dict[str, DeliveredClient] = {}
        for row_number, row in enumerate(rows, start=3):
            if not any(value is not None for value in row):
                continue
            def cell(name: str) -> str:
                value = row[indexes[name]]
                return str(value).strip() if value is not None else ""
            client_id = cell("client_id")
            if not re.fullmatch(r"[A-Za-z0-9_-]+", client_id):
                raise ValueError(f"Missing or unsafe client ID in workbook row {row_number}")
            if client_id in clients:
                raise ValueError(f"Client {client_id} occurs more than once in the delivered workbook")
            phone = cell("phone")
            clients[client_id] = DeliveredClient(client_id, phone, cell("funnel"), normalize_phones(phone))
        if not clients:
            raise ValueError(f"Client workbook has no data rows: {path}")
        return clients
    finally:
        workbook.close()


def select_eligible_photos(
    metadata: dict[str, PhotoMeta], clients: dict[str, DeliveredClient]
) -> tuple[dict[str, PhotoMeta], list[dict[str, object]]]:
    """Intersect source photos with eligible delivered clients and report omissions."""

    selected: dict[str, PhotoMeta] = {}
    exclusions: list[dict[str, object]] = []
    for client_id, client in clients.items():
        if client.funnel not in ALLOWED_FUNNELS:
            reason = "outside_requested_funnels"
        elif not client.normalized_phones:
            reason = "invalid_phone" if re.search(r"\d", client.raw_phone) else "missing_phone"
        elif client_id not in metadata:
            reason = "missing_or_unsupported_photo"
        else:
            # Use only phone numbers actually delivered to Fitbase, which can
            # differ from the old raw client-card field selected alongside BLOBs.
            selected[client_id] = replace(metadata[client_id], normalized_phones=client.normalized_phones)
            continue
        exclusions.append({"client_id": client_id, "funnel": client.funnel,
                           "exported_phone": report_text(client.raw_phone), "reason": reason})
    for client_id in sorted(set(metadata) - set(clients)):
        exclusions.append({"client_id": client_id, "funnel": "", "exported_phone": "",
                           "reason": "not_in_delivered_workbook"})
    return selected, exclusions


def assign_names(metadata: dict[str, PhotoMeta]) -> dict[str, NameAssignment]:
    """Maximize exact-phone filenames; preserve shared-phone images with ID suffixes."""

    if any(not item.normalized_phones for item in metadata.values()):
        raise ValueError("Only valid-phone clients may receive photo filenames")
    sys.setrecursionlimit(max(50_000, len(metadata) * 3))
    phone_owner: dict[str, str] = {}

    def augment(client_id: str, visited_phones: set[str]) -> bool:
        for phone in metadata[client_id].normalized_phones:
            if not re.fullmatch(r"7[0-9]{10}", phone):
                raise ValueError(f"Invalid normalized photo phone for client {client_id}")
            if phone in visited_phones:
                continue
            visited_phones.add(phone)
            previous = phone_owner.get(phone)
            if previous is None or augment(previous, visited_phones):
                phone_owner[phone] = client_id
                return True
        return False

    ordered_clients = sorted(metadata, key=lambda client_id: (
        len(metadata[client_id].normalized_phones),
        int(client_id) if client_id.isdigit() else 10**18, client_id,
    ))
    for client_id in ordered_clients:
        if not re.fullmatch(r"[A-Za-z0-9_-]+", client_id):
            raise ValueError(f"Unsafe client ID for photo filename: {client_id!r}")
        augment(client_id, set())

    matched = {client_id: phone for phone, client_id in phone_owner.items()}
    assignments: dict[str, NameAssignment] = {}
    for client_id, item in metadata.items():
        assigned_phone = matched.get(client_id, "")
        if assigned_phone:
            status = "exact_primary_phone" if assigned_phone == item.normalized_phones[0] else "exact_alternate_phone"
            basename = assigned_phone
        else:
            assigned_phone = item.normalized_phones[0]
            status = "duplicate_phone_client_id_suffix"
            basename = f"{assigned_phone}__{client_id}"
        assignments[client_id] = NameAssignment(basename, assigned_phone, status)
    basenames = [item.basename for item in assignments.values()]
    if len(basenames) != len(set(basenames)):
        raise RuntimeError("Phone filename assignment still contains collisions")
    return assignments


def decode_1c_photo(hex_value: bytes) -> tuple[bytes, str]:
    """Decode one ``_InfoRg6767._Fld6769`` 1C container."""

    try:
        wrapped = bytes.fromhex(hex_value.decode("ascii"))
    except (UnicodeDecodeError, ValueError) as exc:
        raise ValueError("Photo BLOB is not valid hexadecimal BCP output") from exc
    if len(wrapped) < 19:
        raise ValueError(f"Photo BLOB is too short: {len(wrapped)} bytes")

    try:
        unpacked = zlib.decompress(wrapped[18:], -15)
    except zlib.error as exc:
        raise ValueError("Cannot inflate the 1C photo container") from exc
    if len(unpacked) < 20:
        raise ValueError(f"Inflated 1C photo container is too short: {len(unpacked)} bytes")

    metadata_length = struct.unpack("<Q", unpacked[:8])[0]
    payload_header = 8 + metadata_length
    if payload_header + 12 > len(unpacked):
        raise ValueError("1C metadata length points outside the inflated container")
    payload_length = struct.unpack("<Q", unpacked[payload_header + 4 : payload_header + 12])[0]
    payload = unpacked[payload_header + 12 : payload_header + 12 + payload_length]
    if len(payload) != payload_length:
        raise ValueError(f"Photo payload is truncated: expected {payload_length}, got {len(payload)}")

    for signature, extension in IMAGE_SIGNATURES:
        if payload.startswith(signature):
            return payload, extension
    raise ValueError(f"Unknown decoded image signature: {payload[:16].hex()}")


def jpeg_payload(payload: bytes, detected_extension: str) -> bytes:
    """Validate pixels and convert other source formats into RGB JPEG on white."""

    with Image.open(io.BytesIO(payload)) as source:
        source.load()  # Reject a truncated image before publishing the archive.
        if detected_extension == "jpg" and source.format == "JPEG" and source.mode in {"RGB", "L"}:
            return payload
        oriented = ImageOps.exif_transpose(source)
        rgba = oriented.convert("RGBA")
        background = Image.new("RGB", rgba.size, "white")
        background.paste(rgba, mask=rgba.getchannel("A"))
        buffer = io.BytesIO()
        background.save(buffer, format="JPEG", quality=95, subsampling=0)
        return buffer.getvalue()


def zip_info(path: str, cutoff_at: str, *, directory: bool = False) -> zipfile.ZipInfo:
    """Use the run's effective date for deterministic ZIP metadata."""

    timestamp = datetime.strptime(cutoff_at, "%Y-%m-%d %H:%M:%S")
    if not 1980 <= timestamp.year <= 2107:
        raise ValueError("Effective cutoff year is outside the ZIP timestamp range")
    zip_timestamp = (timestamp.year, timestamp.month, timestamp.day, timestamp.hour,
                     timestamp.minute, timestamp.second // 2 * 2)
    normalized = path.rstrip("/") + "/" if directory else path
    info = zipfile.ZipInfo(normalized, zip_timestamp)
    info.create_system = 3
    info.external_attr = (0o40755 << 16) | 0x10 if directory else 0o100644 << 16
    info.compress_type = zipfile.ZIP_STORED if directory else zipfile.ZIP_DEFLATED
    return info


def write_csv_bytes(rows: Sequence[dict[str, object]], fieldnames: Sequence[str]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=fieldnames, extrasaction="ignore", lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")


def write_archive(
    *,
    container: str,
    query: str,
    metadata: dict[str, PhotoMeta],
    assignments: dict[str, NameAssignment],
    clients: dict[str, DeliveredClient],
    exclusions: list[dict[str, object]],
    output_path: Path,
    inner_dir: str,
    cutoff_at: str,
    backup_finish_at: str,
    expected_count: int,
    overwrite: bool,
) -> dict[str, object]:
    """Stream eligible BLOBs into a ZIP, fully validate it, then publish atomically."""

    if not re.fullmatch(r"[A-Za-z0-9_.-]+", inner_dir) or inner_dir in {".", ".."}:
        raise ValueError(f"Unsafe ZIP inner directory: {inner_dir!r}")
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"Output already exists (use --overwrite intentionally): {output_path}")
    if expected_count != len(metadata) or set(metadata) != set(assignments):
        raise ValueError("Expected photo count and assigned IDs must match selected metadata")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = output_path.with_name(output_path.name + ".partial")
    if temp_path.exists():
        raise FileExistsError(f"Unfinished photo export exists: {temp_path}; inspect it before retrying")

    process = start_bcp(container, query)
    unexpected: list[bytes] = []
    manifest_rows: list[dict[str, object]] = []
    seen_clients: set[str] = set()
    seen_paths: set[str] = set()
    status_counts: Counter[str] = Counter()
    source_format_counts: Counter[str] = Counter()
    source_counts: Counter[str] = Counter()
    total_payload_bytes = 0
    manifest_fields = ["filename", "client_id", "funnel", "assigned_phone", "exported_phone",
                       "raw_phone", "naming_status", "photo_source", "metadata_extension",
                       "detected_extension", "output_extension", "bytes", "sha256"]

    try:
        with zipfile.ZipFile(temp_path, mode="w", compression=zipfile.ZIP_DEFLATED,
                             compresslevel=6, allowZip64=True, strict_timestamps=True) as archive:
            for directory in (inner_dir, f"{inner_dir}/photos", f"{inner_dir}/_reports"):
                archive.writestr(zip_info(directory, cutoff_at, directory=True), b"")
            assert process.stdout is not None
            for raw_line in process.stdout:
                line = raw_line.rstrip(b"\r\n")
                fields = line.split(b"\t", 4)
                if len(fields) != 5:
                    if not is_bcp_status_line(line):
                        unexpected.append(line[:500])
                    continue
                client_id = fields[0].decode("utf-8", errors="strict")
                item = metadata.get(client_id)
                if item is None:
                    continue  # Explicitly excluded in the metadata/workbook pass.
                if client_id in seen_clients:
                    raise RuntimeError(f"Duplicate photo row for client {client_id}")
                raw_phone, source, metadata_extension = (
                    field.decode("utf-8", errors="strict") for field in fields[1:4])
                if (raw_phone, source, metadata_extension) != (item.raw_phone, item.source, item.metadata_extension):
                    raise RuntimeError(f"Metadata changed between BCP passes for client {client_id}")
                assignment = assignments[client_id]
                payload, detected_extension = decode_1c_photo(fields[4])
                try:
                    payload = jpeg_payload(payload, detected_extension)
                except Exception as exc:
                    raise ValueError(f"Photo cannot be decoded as a complete image: client {client_id}") from exc
                relative_path = f"{inner_dir}/photos/{assignment.basename}.jpg"
                if relative_path in seen_paths:
                    raise RuntimeError(f"Duplicate archive path: {relative_path}")
                archive.writestr(zip_info(relative_path, cutoff_at), payload)
                client = clients[client_id]
                manifest_rows.append({
                    "filename": f"{assignment.basename}.jpg", "client_id": client_id,
                    "funnel": client.funnel, "assigned_phone": assignment.assigned_phone,
                    "exported_phone": report_text(client.raw_phone), "raw_phone": report_text(raw_phone),
                    "naming_status": assignment.status, "photo_source": source,
                    "metadata_extension": metadata_extension, "detected_extension": detected_extension,
                    "output_extension": "jpg", "bytes": len(payload),
                    "sha256": hashlib.sha256(payload).hexdigest(),
                })
                seen_clients.add(client_id)
                seen_paths.add(relative_path)
                status_counts[assignment.status] += 1
                source_format_counts[detected_extension] += 1
                source_counts[source] += 1
                total_payload_bytes += len(payload)
            finish_bcp(process, unexpected)
            if seen_clients != set(metadata):
                missing = sorted(set(metadata) - seen_clients)
                raise RuntimeError(f"Photo/metadata client mismatch: missing={missing[:10]}")
            if len(seen_clients) != expected_count:
                raise RuntimeError(f"Expected {expected_count} photos, wrote {len(seen_clients)}")

            manifest_rows.sort(key=lambda row: (str(row["client_id"]), str(row["filename"])))
            exceptions = [row for row in manifest_rows if not str(row["naming_status"]).startswith("exact_")]
            for filename, rows, fields in (
                ("manifest.csv", manifest_rows, manifest_fields),
                ("phone_filename_exceptions.csv", exceptions, manifest_fields),
                ("excluded_clients.csv", exclusions, ["client_id", "funnel", "exported_phone", "reason"]),
            ):
                archive.writestr(zip_info(f"{inner_dir}/_reports/{filename}", cutoff_at), write_csv_bytes(rows, fields))
            readme = f"""Фотографии действующих клиентов и реактивации из итогового XLSX

Фактическое завершение backup: {backup_finish_at}
Актуальность данных (cutoff_at): {cutoff_at}
Фотографий: {len(manifest_rows)}

Каждая фотография — проверенный JPEG. Телефон взят из итогового файла заявок.
Обычное имя: <нормализованный телефон>.jpg
Общий валидный телефон: <телефон>__<ID клиента>.jpg
Невалидные и отсутствующие телефоны, отсутствующие фото исключены из архива.
На состав основных XLSX это исключение не влияет.

Соответствие клиентам и SHA-256: _reports/manifest.csv
Общие телефоны с суффиксом: _reports/phone_filename_exceptions.csv
Исключённые из фотоархива клиенты: _reports/excluded_clients.csv
"""
            archive.writestr(zip_info(f"{inner_dir}/README.txt", cutoff_at), readme.encode("utf-8"))
        validation = validate_archive(temp_path, inner_dir, expected_count)
        if output_path.exists() and not overwrite:
            raise FileExistsError(f"Output appeared while exporting: {output_path}")
        temp_path.replace(output_path)
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise
    finally:
        close_bcp(process)
    return {
        "photos": len(seen_clients), "payload_bytes": total_payload_bytes,
        "format_counts": {"jpg": len(seen_clients)},
        "source_format_counts": dict(sorted(source_format_counts.items())),
        "source_counts": dict(sorted(source_counts.items())),
        "naming_status_counts": dict(sorted(status_counts.items())),
        "phone_filename_exceptions": len(exceptions),
        "excluded_clients": len(exclusions),
        "exclusion_counts": dict(sorted(Counter(str(row["reason"]) for row in exclusions).items())),
        "funnel_counts": dict(sorted(Counter(str(row["funnel"]) for row in manifest_rows).items())),
        "validation": validation,
    }


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_archive(path: Path, inner_dir: str, expected_count: int) -> dict[str, object]:
    """Read back CRC, full JPEG pixels and exact manifest membership/hashes."""

    photo_prefix = f"{inner_dir}/photos/"
    with zipfile.ZipFile(path, "r") as archive:
        bad_crc = archive.testzip()
        if bad_crc is not None:
            raise RuntimeError(f"ZIP CRC validation failed at {bad_crc}")
        names = archive.namelist()
        if len(names) != len(set(names)):
            raise RuntimeError("ZIP contains duplicate entry paths")
        photo_infos = [info for info in archive.infolist()
                       if info.filename.startswith(photo_prefix) and not info.is_dir()]
        if len(photo_infos) != expected_count:
            raise RuntimeError(f"ZIP contains {len(photo_infos)} photos, expected {expected_count}")
        with archive.open(f"{inner_dir}/_reports/manifest.csv", "r") as handle:
            rows = list(csv.DictReader(io.TextIOWrapper(handle, encoding="utf-8-sig", newline="")))
        manifest = {row["filename"]: row for row in rows}
        if len(rows) != expected_count or len(manifest) != expected_count:
            raise RuntimeError("Photo manifest row count or unique filename count mismatch")
        if len({row["client_id"] for row in rows}) != expected_count:
            raise RuntimeError("Photo manifest has duplicate client IDs")
        if set(manifest) != {info.filename.removeprefix(photo_prefix) for info in photo_infos}:
            raise RuntimeError("Photo manifest and archive entry names differ")
        for info in photo_infos:
            filename = info.filename.removeprefix(photo_prefix)
            row = manifest[filename]
            if not re.fullmatch(r"7[0-9]{10}(?:__[A-Za-z0-9_-]+)?\.jpg", filename):
                raise RuntimeError(f"Photo filename is not based on a valid phone: {filename}")
            if filename.split("__", 1)[0].removesuffix(".jpg") != row["assigned_phone"]:
                raise RuntimeError(f"Photo filename and assigned phone differ: {filename}")
            if row["assigned_phone"] not in normalize_phones(row["exported_phone"]):
                raise RuntimeError(f"Photo phone is absent from the delivered client row: {filename}")
            if row["funnel"] not in ALLOWED_FUNNELS:
                raise RuntimeError(f"Unexpected photo funnel: {filename}")
            payload = archive.read(info)
            if int(row["bytes"]) != len(payload) or row["sha256"] != hashlib.sha256(payload).hexdigest():
                raise RuntimeError(f"Photo byte count or SHA-256 differs from manifest: {filename}")
            with Image.open(io.BytesIO(payload)) as decoded:
                if decoded.format != "JPEG":
                    raise RuntimeError(f"Photo is not JPEG: {filename}")
                decoded.load()
        total_uncompressed = sum(info.file_size for info in photo_infos)
        total_compressed = sum(info.compress_size for info in photo_infos)
    return {
        "zip_bytes": path.stat().st_size, "zip_sha256": sha256_file(path),
        "photo_entries": len(photo_infos), "manifest_rows": len(rows),
        "photo_uncompressed_bytes": total_uncompressed, "photo_compressed_bytes": total_compressed,
        "signature_counts": {"jpg": len(photo_infos)}, "crc_check": "PASS",
        "unique_paths_check": "PASS", "decoded_images_check": "PASS", "manifest_hashes_check": "PASS",
    }


def run_preflight(database: str, container: str, cutoff_at: str, backup_finish_at: str) -> None:
    """Require every built SQL export layer to use the effective business cutoff."""

    cutoff = datetime.strptime(cutoff_at, "%Y-%m-%d %H:%M:%S")
    cutoff_date = cutoff.strftime("%Y%m%d")
    datetime.strptime(backup_finish_at, "%Y-%m-%d %H:%M:%S")
    db = require_safe_identifier(database, "database name")
    require_safe_container(container)
    sql = " ".join(f"""
        SET NOCOUNT ON;
        USE [{db}];
        DECLARE @cutoff_at datetime2(0) = CONVERT(datetime2(0), '{cutoff_at}', 120);
        DECLARE @cutoff_date date = CONVERT(date, '{cutoff_date}', 112);
        DECLARE @backup_finish_at datetime2(0) = CONVERT(datetime2(0), '{backup_finish_at}', 120);
        IF DB_NAME() <> N'{db}' THROW 51000, 'Unexpected database', 1;
        IF (SELECT COUNT(*) FROM fitbase_part2.staging_run_metadata) <> 1
            THROW 51005, 'Expected exactly one owner staging metadata row', 1;
        IF EXISTS (SELECT 1 FROM fitbase_part2.staging_run_metadata
                   WHERE cutoff_date <> @cutoff_date OR cutoff_date IS NULL
                      OR cutoff_at <> @cutoff_at OR cutoff_at IS NULL
                      OR backup_finish_at <> @backup_finish_at OR backup_finish_at IS NULL)
            THROW 51006, 'Owner effective cutoff or actual backup provenance mismatch', 1;
        IF EXISTS (SELECT 1 FROM fitbase_part2.final_funnel_clients
                   WHERE cutoff_date <> @cutoff_date OR cutoff_date IS NULL)
            THROW 51001, 'final_funnel_clients cutoff mismatch', 1;
        IF EXISTS (SELECT 1 FROM fitbase_part2.membership_import_facts
                   WHERE cutoff_at <> @cutoff_at OR cutoff_at IS NULL)
            THROW 51002, 'membership cutoff mismatch', 1;
        IF EXISTS (SELECT 1 FROM fitbase_part2.services_import_facts
                   WHERE cutoff_at <> @cutoff_at OR cutoff_at IS NULL)
            THROW 51003, 'services cutoff mismatch', 1;
        IF NOT EXISTS (SELECT 1 FROM fitbase_part2.final_funnel_clients WHERE cutoff_date = @cutoff_date)
            THROW 51004, 'No client rows at requested cutoff', 1;
        SELECT 'cutoff_contract' AS check_name, 'PASS' AS result,
               @@SERVERNAME AS server_name, DB_NAME() AS database_name,
               @cutoff_at AS cutoff_at, @backup_finish_at AS backup_finish_at;
        """.split())
    subprocess.run([
        "docker", "exec", container, "/bin/bash", "-lc",
        'exec /opt/mssql-tools18/bin/sqlcmd -S localhost -U sa -P "$MSSQL_SA_PASSWORD" -C "$@"',
        "sqlcmd", "-d", db, "-b", "-Q", sql,
    ], cwd=ROOT, check=True)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--container", required=True)
    parser.add_argument("--database", required=True)
    parser.add_argument("--cutoff-at", required=True, help="Effective business timestamp YYYY-MM-DD HH:MM:SS")
    parser.add_argument("--backup-finish-at", required=True,
                        help="Actual RESTORE HEADERONLY.BackupFinishDate already verified by the pipeline")
    parser.add_argument("--clients-xlsx", type=Path, required=True, help="Final combined client workbook to be delivered")
    parser.add_argument("--expected-photo-count", type=int, default=None,
                        help="Optional rehearsal/reference assertion; no historical count is assumed")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--inner-dir", default=None)
    parser.add_argument("--report-json", type=Path, default=None)
    parser.add_argument("--overwrite", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    cutoff = datetime.strptime(args.cutoff_at, "%Y-%m-%d %H:%M:%S")
    datetime.strptime(args.backup_finish_at, "%Y-%m-%d %H:%M:%S")
    cutoff_date = cutoff.strftime("%Y%m%d")
    inner_dir = args.inner_dir or f"fitbase_client_photos_{cutoff_date}"
    output_path = args.output.resolve()
    clients_path = args.clients_xlsx.resolve()
    if not re.fullmatch(r"[A-Za-z0-9_.-]+", inner_dir) or inner_dir in {".", ".."}:
        raise ValueError(f"Unsafe ZIP inner directory: {inner_dir!r}")
    if args.expected_photo_count is not None and args.expected_photo_count < 0:
        raise ValueError("Expected photo count must not be negative")
    clients = load_delivered_clients(clients_path)
    run_preflight(args.database, args.container, args.cutoff_at, args.backup_finish_at)
    metadata_query = sql_selected_union(args.database, include_blob=False, cutoff_date=cutoff_date)
    photo_query = sql_selected_union(args.database, include_blob=True, cutoff_date=cutoff_date)
    source_metadata = load_metadata(args.container, metadata_query)
    metadata, exclusions = select_eligible_photos(source_metadata, clients)
    if not metadata:
        raise RuntimeError("No eligible client photos found; inspect workbook scope and photo sources")
    if args.expected_photo_count is not None and len(metadata) != args.expected_photo_count:
        raise RuntimeError(f"Expected {args.expected_photo_count} eligible photos, found {len(metadata)}")
    build_summary = write_archive(
        container=args.container, query=photo_query, metadata=metadata, assignments=assign_names(metadata),
        clients=clients, exclusions=exclusions, output_path=output_path, inner_dir=inner_dir,
        cutoff_at=args.cutoff_at, backup_finish_at=args.backup_finish_at,
        expected_count=len(metadata), overwrite=args.overwrite,
    )
    result = {
        "status": "PASS", "output": str(output_path), "inner_directory": inner_dir,
        "database": args.database, "container": args.container,
        "cutoff_at": args.cutoff_at, "backup_finish_at": args.backup_finish_at,
        "clients_xlsx": str(clients_path), "clients_xlsx_sha256": sha256_file(clients_path),
        "delivered_clients": len(clients), "source_photo_candidates": len(source_metadata),
        "build": build_summary, "validation": build_summary.pop("validation"),
    }
    report = json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    if args.report_json:
        args.report_json.parent.mkdir(parents=True, exist_ok=True)
        args.report_json.write_text(report, encoding="utf-8")
    print(report, end="")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
