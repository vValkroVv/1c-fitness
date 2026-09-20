#!/usr/bin/env python3
"""Build a delivery where only the final reactivation labels may change."""

from __future__ import annotations

import argparse
import csv
import hashlib
import shutil
from pathlib import Path
from typing import Any
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OWNER_PREFIX = "fitbase_active_clients_import_zayavki_"
OLD_FUNNEL = "Реактивация(годовые абонементы)"
NEW_FUNNEL = "Реактивация"
OLD_STEP = "Все закрытые абонементы"
NEW_STEP = "Закрытые годовые абонементы"
CORE_PROPERTIES = "docProps/core.xml"
WORKSHEET_XML = "xl/worksheets/sheet1.xml"
SHEET_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def as_abs(path: str | Path) -> Path:
    value = Path(path)
    return value.resolve() if value.is_absolute() else (ROOT / value).resolve()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def blank(value: Any) -> bool:
    return value in (None, "")


def read_rows(path: Path) -> tuple[list[str], list[str], list[tuple[Any, ...]], str]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    if len(workbook.sheetnames) != 1:
        workbook.close()
        raise ValueError(f"Expected exactly one sheet in {path}")
    worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    headers = [str(value or "") for value in next(iterator)]
    russian_headers = [str(value or "") for value in next(iterator)]
    rows = [
        tuple(values[: len(headers)])
        for values in iterator
        if any(not blank(value) for value in values[: len(headers)])
    ]
    title = worksheet.title
    workbook.close()
    return headers, russian_headers, rows, title


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def cell_structure(cell: ElementTree.Element) -> list[tuple[str, tuple[tuple[str, str], ...], str]]:
    """Return OOXML structure while ignoring only the stored cell value text."""

    result: list[tuple[str, tuple[tuple[str, str], ...], str]] = []
    for element in cell.iter():
        name = local_name(element.tag)
        text = "" if name in {"t", "v"} else (element.text or "")
        result.append((name, tuple(sorted(element.attrib.items())), text))
    return result


def validate_ooxml_structure(
    old_path: Path,
    corrected_path: Path,
    label_columns: set[str],
) -> None:
    """Verify styles/layout without slow per-cell openpyxl object access.

    The builder writes inline strings, so only the worksheet XML and the normal
    modified timestamp may differ. Every other OOXML member must stay byte-for-
    byte identical. Inside the worksheet, only the text payload of data cells in
    the resolved funnel/funnel_step columns may differ; cell attributes, style
    IDs, formulas, row metadata, filters, panes, dimensions and all other
    worksheet structures must match.
    """

    with ZipFile(old_path) as old_zip, ZipFile(corrected_path) as new_zip:
        old_names = set(old_zip.namelist())
        new_names = set(new_zip.namelist())
        if old_names != new_names:
            raise ValueError("Corrected owner workbook changed OOXML member set")
        allowed_different_members = {CORE_PROPERTIES, WORKSHEET_XML}
        for name in sorted(old_names - allowed_different_members):
            if old_zip.read(name) != new_zip.read(name):
                raise ValueError(
                    f"Corrected owner workbook changed protected OOXML member: {name}"
                )
        old_core = ElementTree.fromstring(old_zip.read(CORE_PROPERTIES))
        new_core = ElementTree.fromstring(new_zip.read(CORE_PROPERTIES))
        for element in old_core.iter():
            if local_name(element.tag) == "modified":
                element.text = ""
        for element in new_core.iter():
            if local_name(element.tag) == "modified":
                element.text = ""
        if ElementTree.tostring(old_core) != ElementTree.tostring(new_core):
            raise ValueError(
                "Corrected owner workbook changed core metadata beyond modified timestamp"
            )
        old_root = ElementTree.fromstring(old_zip.read(WORKSHEET_XML))
        new_root = ElementTree.fromstring(new_zip.read(WORKSHEET_XML))

    if old_root.attrib != new_root.attrib:
        raise ValueError("Corrected owner workbook changed worksheet root attributes")
    old_children = list(old_root)
    new_children = list(new_root)
    if [child.tag for child in old_children] != [child.tag for child in new_children]:
        raise ValueError("Corrected owner workbook changed worksheet child topology")

    sheet_data_tag = f"{{{SHEET_MAIN_NS}}}sheetData"
    row_tag = f"{{{SHEET_MAIN_NS}}}row"
    cell_tag = f"{{{SHEET_MAIN_NS}}}c"
    for old_child, new_child in zip(old_children, new_children, strict=True):
        if old_child.tag != sheet_data_tag:
            if ElementTree.tostring(old_child) != ElementTree.tostring(new_child):
                raise ValueError(
                    "Corrected owner workbook changed worksheet layout element: "
                    f"{local_name(old_child.tag)}"
                )
            continue

        old_rows = old_child.findall(row_tag)
        new_rows = new_child.findall(row_tag)
        if len(old_rows) != len(new_rows):
            raise ValueError("Corrected owner workbook changed physical row count")
        for old_row, new_row in zip(old_rows, new_rows, strict=True):
            if old_row.attrib != new_row.attrib:
                raise ValueError(
                    f"Corrected owner workbook changed row metadata: {old_row.attrib.get('r')}"
                )
            old_cells = old_row.findall(cell_tag)
            new_cells = new_row.findall(cell_tag)
            if len(old_cells) != len(new_cells):
                raise ValueError(
                    f"Corrected owner workbook changed cell topology at row {old_row.attrib.get('r')}"
                )
            for old_cell, new_cell in zip(old_cells, new_cells, strict=True):
                coordinate = old_cell.attrib.get("r", "")
                if old_cell.attrib != new_cell.attrib:
                    raise ValueError(
                        f"Corrected owner workbook changed cell attributes: {coordinate}"
                    )
                row_number = int(old_row.attrib.get("r", "0"))
                column = "".join(character for character in coordinate if character.isalpha())
                is_label_cell = row_number >= 3 and column in label_columns
                if is_label_cell:
                    if cell_structure(old_cell) != cell_structure(new_cell):
                        raise ValueError(
                            f"Corrected owner workbook changed label-cell structure: {coordinate}"
                        )
                elif ElementTree.tostring(old_cell) != ElementTree.tostring(new_cell):
                    raise ValueError(
                        f"Corrected owner workbook changed protected cell XML: {coordinate}"
                    )


def validate_targeted_change(old_path: Path, corrected_path: Path) -> dict[str, int]:
    old_headers, old_russian, old_rows, old_sheet = read_rows(old_path)
    new_headers, new_russian, new_rows, new_sheet = read_rows(corrected_path)
    if old_headers != new_headers:
        raise ValueError("Corrected owner workbook changed technical headers")
    if old_russian != new_russian:
        raise ValueError("Corrected owner workbook changed Russian headers")
    if old_sheet != new_sheet:
        raise ValueError("Corrected owner workbook changed sheet name")
    if len(old_rows) != len(new_rows):
        raise ValueError(
            "Corrected owner workbook changed row count: "
            f"{len(old_rows)} -> {len(new_rows)}"
        )

    indexes = {name: index for index, name in enumerate(old_headers)}
    required = {"client_id", "funnel", "funnel_step"}
    missing = sorted(required - set(indexes))
    if missing:
        raise ValueError(f"Missing owner workbook columns: {missing}")
    client_index = indexes["client_id"]
    funnel_index = indexes["funnel"]
    step_index = indexes["funnel_step"]
    allowed_indexes = {funnel_index, step_index}

    changed_cells = 0
    changed_rows = 0
    reactivation_rows = 0
    for row_number, (old_row, new_row) in enumerate(
        zip(old_rows, new_rows, strict=True),
        start=3,
    ):
        old_id = str(old_row[client_index] or "").strip()
        new_id = str(new_row[client_index] or "").strip()
        if old_id != new_id:
            raise ValueError(
                f"Client identity/order changed at Excel row {row_number}: "
                f"{old_id!r} -> {new_id!r}"
            )

        old_pair = (old_row[funnel_index], old_row[step_index])
        new_pair = (new_row[funnel_index], new_row[step_index])
        is_reactivation = old_pair == (OLD_FUNNEL, OLD_STEP)
        if is_reactivation:
            reactivation_rows += 1
            if new_pair != (NEW_FUNNEL, NEW_STEP):
                raise ValueError(
                    "Reactivation labels were not replaced exactly: "
                    f"row={row_number}, old={old_pair!r}, new={new_pair!r}"
                )
        elif OLD_FUNNEL in old_pair or OLD_STEP in old_pair:
            raise ValueError(
                f"Source owner workbook has a partial legacy pair at row {row_number}: "
                f"{old_pair!r}"
            )
        elif new_pair != old_pair:
            raise ValueError(
                f"Non-reactivation funnel labels changed at row {row_number}: "
                f"{old_pair!r} -> {new_pair!r}"
            )

        row_changes = 0
        for index, (old_value, new_value) in enumerate(
            zip(old_row, new_row, strict=True)
        ):
            if old_value == new_value:
                continue
            if index not in allowed_indexes:
                raise ValueError(
                    "Corrected owner workbook changed a non-authorized cell: "
                    f"row={row_number}, field={old_headers[index]!r}, "
                    f"old={old_value!r}, new={new_value!r}"
                )
            changed_cells += 1
            row_changes += 1
        if row_changes:
            changed_rows += 1
        if is_reactivation and row_changes != 2:
            raise ValueError(
                f"Expected exactly two changed label cells at row {row_number}, "
                f"found {row_changes}"
            )

    if reactivation_rows == 0:
        raise ValueError("Source owner workbook contains no legacy reactivation rows")
    if changed_rows != reactivation_rows:
        raise ValueError(
            "Changed-row count does not equal reactivation-row count: "
            f"{changed_rows} != {reactivation_rows}"
        )
    if changed_cells != reactivation_rows * 2:
        raise ValueError(
            "Changed-cell count does not equal two labels per reactivation row: "
            f"{changed_cells} != {reactivation_rows * 2}"
        )
    if any(
        row[funnel_index] == OLD_FUNNEL or row[step_index] == OLD_STEP
        for row in new_rows
    ):
        raise ValueError("Legacy reactivation labels remain in corrected workbook")

    def column_letter(index: int) -> str:
        number = index + 1
        letters = ""
        while number:
            number, remainder = divmod(number - 1, 26)
            letters = chr(65 + remainder) + letters
        return letters

    validate_ooxml_structure(
        old_path,
        corrected_path,
        {column_letter(funnel_index), column_letter(step_index)},
    )

    return {
        "rows": len(new_rows),
        "reactivation_rows": reactivation_rows,
        "changed_rows": changed_rows,
        "changed_cells": changed_cells,
    }


def write_manifest(path: Path, rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "file_name",
                "source_sha256",
                "delivery_sha256",
                "status",
            ],
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(rows)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-delivery", required=True)
    parser.add_argument("--corrected-owner", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--date-stamp", default="20260630")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    source_dir = as_abs(args.source_delivery)
    corrected_owner = as_abs(args.corrected_owner)
    output_dir = as_abs(args.output_dir)
    owner_name = f"{OWNER_PREFIX}{args.date_stamp}_all_funnels.xlsx"
    old_owner = source_dir / owner_name

    if not source_dir.is_dir():
        raise FileNotFoundError(source_dir)
    if not old_owner.is_file():
        raise FileNotFoundError(old_owner)
    if not corrected_owner.is_file():
        raise FileNotFoundError(corrected_owner)
    if corrected_owner.name != owner_name:
        raise ValueError(
            f"Corrected owner file must be named {owner_name!r}, "
            f"got {corrected_owner.name!r}"
        )
    if output_dir.exists() and any(output_dir.iterdir()):
        raise FileExistsError(
            f"Targeted delivery already exists and is not empty: {output_dir}"
        )

    change_validation = validate_targeted_change(old_owner, corrected_owner)

    if output_dir.exists():
        output_dir.rmdir()
    shutil.copytree(
        source_dir,
        output_dir,
        copy_function=shutil.copy2,
        ignore=shutil.ignore_patterns(".DS_Store"),
    )
    shutil.copy2(corrected_owner, output_dir / owner_name)

    source_xlsx = sorted(source_dir.rglob("*.xlsx"))
    manifest_rows: list[dict[str, str]] = []
    for source_path in source_xlsx:
        relative_path = source_path.relative_to(source_dir)
        destination = output_dir / relative_path
        source_hash = sha256(source_path)
        delivery_hash = sha256(destination)
        is_owner = relative_path == Path(owner_name)
        if is_owner and source_hash == delivery_hash:
            raise ValueError("Corrected owner workbook is byte-identical to the old one")
        if not is_owner and source_hash != delivery_hash:
            raise ValueError(
                f"Untargeted workbook changed while copying: {relative_path}"
            )
        manifest_rows.append(
            {
                "file_name": relative_path.as_posix(),
                "source_sha256": source_hash,
                "delivery_sha256": delivery_hash,
                "status": "changed_funnel_labels_only" if is_owner else "byte_identical",
            }
        )

    reports_dir = output_dir / "reports"
    reports_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = reports_dir / "funnel_label_xlsx_hash_manifest.csv"
    write_manifest(manifest_path, manifest_rows)

    root_xlsx_count = sum(path.parent == source_dir for path in source_xlsx)
    unchanged_root_xlsx = root_xlsx_count - 1
    unchanged_recursive_xlsx = len(manifest_rows) - 1
    report = [
        "# Targeted funnel-label delivery",
        "",
        f"- source delivery: `{source_dir}`",
        f"- corrected delivery: `{output_dir}`",
        f"- total owner rows: {change_validation['rows']}",
        f"- reactivation rows: {change_validation['reactivation_rows']}",
        f"- rows with changed labels: {change_validation['changed_rows']}",
        f"- changed cells (funnel and funnel_step only): {change_validation['changed_cells']}",
        f"- unchanged root XLSX verified byte-identical: {unchanged_root_xlsx}",
        f"- unchanged recursive XLSX verified byte-identical: {unchanged_recursive_xlsx}",
        f"- old funnel: `{OLD_FUNNEL}`",
        f"- new funnel: `{NEW_FUNNEL}`",
        f"- old funnel_step: `{OLD_STEP}`",
        f"- new funnel_step: `{NEW_STEP}`",
        "- worksheet layout, cell attributes and every style ID: identical",
        "- protected OOXML members: byte-identical",
        "- core metadata difference: modified timestamp only",
        "- status: PASS",
        "",
        "Only the two requested labels are authorized to differ inside the owner workbook.",
        "Every other cell and every other XLSX is checked before the delivery is accepted.",
    ]
    (reports_dir / "funnel_label_change_validation.md").write_text(
        "\n".join(report) + "\n",
        encoding="utf-8",
    )
    print("\n".join(report))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
