"""Delivery publication and resume barriers, without a live SQL connection."""

from __future__ import annotations

import argparse
import contextlib
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import yaml
from openpyxl import Workbook, load_workbook


SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))
import run_pipeline as runner


class PipelineSafetyTests(unittest.TestCase):
    def setUp(self) -> None:
        directory = self.enterContext(tempfile.TemporaryDirectory())
        self.root = Path(directory)
        self.config_dir = self.root / "config"
        self.config_dir.mkdir()
        self.managers = self.config_dir / "managers_by_club.yml"
        self.managers.write_text("assignment_mode: global\nmanagers: [Manager One]\n", encoding="utf-8")
        (self.config_dir / "branches_by_club.yml").write_text("branches: {}\n", encoding="utf-8")
        (self.config_dir / "product_reclassification_decisions.csv").write_text("product_id,decision\n", encoding="utf-8")
        self.expected = self.config_dir / "expected.yml"
        self.expected.write_text("{}\n", encoding="utf-8")
        self.config = {
            "run": {"work_name": "guardrail_test", "delivery_name": "guardrail_delivery", "effective_offset_days": 1},
            "sql": {"server": "unused.invalid", "port": 1433, "database": "UnusedTestDatabase", "user": "unused",
                    "password_env": "FITBASE_GUARDRAIL_TEST_PASSWORD"},
            "validation": {"expected_manifest": str(self.expected), "enforce_reference_counts": False},
            "backup": {"backup_finish_at": "2026-09-20 21:15:07"},
            "delivery": {"output_base": str(self.root / "output")},
            "photos": {"enabled": False},
        }
        self.config_path = self.config_dir / "pipeline.yml"
        self.write_config()
        self.args = argparse.Namespace(
            config=str(self.config_path), server=None, port=None, database=None, user=None,
            password_env=None, password_file=None, resume=False, skip_reference_counts=False,
            start_at="preflight", stop_after="manifest",
        )
        self.enterContext(patch.multiple(runner, ROOT=self.root, CONFIG=self.config_dir))
        self.enterContext(patch.dict("os.environ", {"FITBASE_GUARDRAIL_TEST_PASSWORD": "unused-test-password"}))
        self.database = self.enterContext(patch.object(runner, "DatabaseClient", side_effect=AssertionError("Test attempted live SQL")))
        self.enterContext(contextlib.redirect_stdout(io.StringIO()))

    def write_config(self) -> None:
        self.config_path.write_text(yaml.safe_dump(self.config, allow_unicode=True), encoding="utf-8")

    def pipeline(self):
        return runner.Pipeline(self.args)

    def test_mutated_xlsx_after_validation_cannot_republish_ready_delivery(self):
        pipeline = self.pipeline()
        pipeline.prepare_directories()
        workbook_path = pipeline.delivery_root / "clients.xlsx"
        workbook = Workbook()
        workbook.active.append(["client_id", "phone"])
        workbook.active.append(["000001", "79991112233"])
        workbook.save(workbook_path)
        workbook.close()

        # Structural validation itself has its own tests. Here its successful
        # command result feeds the real validation fingerprint capture.
        def successful_validation(step, command):
            self.assertEqual(step, "delivery_validate")
            reports = pipeline.delivery_root / "reports"
            (reports / "structural_validation.md").write_text("PASS\n", encoding="utf-8")
            (reports / "structural_validation.json").write_text('{"verdict": "PASS"}\n', encoding="utf-8")

        with patch.object(pipeline, "run_command", side_effect=successful_validation):
            pipeline.validate()
        pipeline.manifest()
        self.assertTrue((pipeline.delivery_root / "READY.txt").is_file())

        changed = load_workbook(workbook_path)
        changed.active["B2"] = "79994445566"
        changed.save(workbook_path)
        changed.close()
        pipeline.status["completed_steps"] = list(runner.STEPS)
        pipeline._write_status()
        self.args.resume = True
        self.args.start_at = "manifest"
        with self.assertRaisesRegex(RuntimeError, "changed since validation"):
            pipeline.run()
        self.assertFalse((pipeline.delivery_root / "READY.txt").exists())
        self.assertFalse((pipeline.delivery_root / "reports" / "delivery_manifest.json").exists())
        self.assertNotIn("manifest", pipeline.status["completed_steps"])
        self.assertIn("failed_at", json.loads(pipeline.status_path.read_text(encoding="utf-8")))
        self.database.assert_not_called()

    def test_resume_rejects_changed_manager_file_or_effective_timestamp(self):
        pipeline = self.pipeline()
        pipeline.prepare_directories()
        self.args.resume = True
        same = self.pipeline()
        self.assertEqual(same.status["config_signature"], pipeline.status["config_signature"])

        original_managers = self.managers.read_bytes()
        self.managers.write_text("assignment_mode: global\nmanagers: [Manager Two]\n", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "Cannot resume with changed configuration or cutoff"):
            self.pipeline()
        self.managers.write_bytes(original_managers)

        self.config["run"].pop("effective_offset_days")
        self.config["run"]["effective_at"] = "2026-09-21 21:15:08"
        self.write_config()
        with self.assertRaisesRegex(ValueError, "Cannot resume with changed configuration or cutoff"):
            self.pipeline()
        self.database.assert_not_called()

    def test_restart_earlier_step_invalidates_later_success_and_publication_markers(self):
        pipeline = self.pipeline()
        pipeline.prepare_directories()
        pipeline.status["completed_steps"] = list(runner.STEPS)
        pipeline._write_status()
        (pipeline.delivery_root / "READY.txt").write_text("old PASS\n", encoding="utf-8")
        reports = pipeline.delivery_root / "reports"
        reports.mkdir()
        (reports / "delivery_manifest.json").write_text('{"verdict":"PASS"}\n', encoding="utf-8")
        self.args.resume = True
        self.args.start_at = "main_xlsx"
        self.args.stop_after = "main_xlsx"
        prior = runner.STEPS[:runner.STEPS.index("main_xlsx")]

        def rebuild_clients():
            self.assertEqual(pipeline.status["completed_steps"], prior)
            self.assertFalse((pipeline.delivery_root / "READY.txt").exists())
            self.assertFalse((reports / "delivery_manifest.json").exists())

        with patch.object(pipeline, "main_xlsx", side_effect=rebuild_clients) as rebuild:
            pipeline.run()
        rebuild.assert_called_once()
        self.assertEqual(pipeline.status["completed_steps"], [*prior, "main_xlsx"])
        saved = json.loads(pipeline.status_path.read_text(encoding="utf-8"))
        self.assertEqual(saved["completed_steps"], [*prior, "main_xlsx"])
        self.assertFalse((pipeline.delivery_root / "READY.txt").exists())
        self.database.assert_not_called()


if __name__ == "__main__":
    unittest.main()
