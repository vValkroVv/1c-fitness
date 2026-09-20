from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
EXPECTED_MANAGERS = {
    "Пеуна Анастасия Ивановна",
    "Пилия Анастасия Артуровна",
    "Ефремова Алена",
}


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


BUILDERS = [
    (ROOT, load_module("packaged_manager_builder", ROOT / "scripts/12_build_part2_three_funnel_xlsx.py")),
]
if (ROOT.parent / "scripts/12_build_part2_three_funnel_xlsx.py").exists():
    BUILDERS.append(
        (ROOT.parent, load_module("root_manager_builder", ROOT.parent / "scripts/12_build_part2_three_funnel_xlsx.py"))
    )
VALIDATORS = [
    (root, load_module(f"manager_validator_{index}", root / "scripts/18_validate_combined_single_stage_outputs.py"))
    for index, (root, _) in enumerate(BUILDERS)
]


def client_rows():
    clubs = ["Коммунальная, 20", "Ровио, 3", "Новый клуб", ""]
    funnels = ["Действующие клиенты", "Реактивация", "Новые заявки"]
    return [
        {
            "client_id": f"{client:09}",
            "normalized_club": club,
            "funnel": funnel,
        }
        for client in range(100)
        for club in clubs
        for funnel in funnels
    ]


class GlobalManagerAssignmentTests(unittest.TestCase):
    def test_default_pool_contains_exactly_requested_three_managers(self):
        for root, builder in BUILDERS:
            with self.subTest(root=root):
                pools = builder.load_managers(root / "config/managers_by_club.yml")
                self.assertEqual(set(pools), {"*"})
                self.assertEqual(len(pools["*"]), 3)
                self.assertEqual(set(pools["*"]), EXPECTED_MANAGERS)

    def test_assignment_is_independent_of_club_funnel_and_row_order(self):
        for root, builder in BUILDERS:
            with self.subTest(root=root):
                pools = builder.load_managers(root / "config/managers_by_club.yml")
                rows = client_rows()
                builder.assign_managers(rows, pools)
                by_client = {}
                for row in rows:
                    by_client.setdefault(row["client_id"], set()).add(row["manager"])
                self.assertTrue(all(len(names) == 1 for names in by_client.values()))
                self.assertEqual({row["manager"] for row in rows}, EXPECTED_MANAGERS)
                reversed_rows = list(reversed(client_rows()))
                builder.assign_managers(reversed_rows, pools)
                self.assertEqual(
                    [(row["client_id"], row["manager"]) for row in rows],
                    [(row["client_id"], row["manager"]) for row in reversed(reversed_rows)],
                )

    def test_legacy_club_pools_remain_supported(self):
        with tempfile.TemporaryDirectory() as directory:
            config = Path(directory) / "legacy.yml"
            config.write_text("clubs:\n  Club A: [Alice]\n  Club B: [Bob]\n", encoding="utf-8")
            for _, builder in BUILDERS:
                pools = builder.load_managers(config)
                rows = [{"client_id": "1", "normalized_club": club} for club in ["Club A", "Club B", "Unknown"]]
                builder.assign_managers(rows, pools)
                self.assertEqual([row["manager"] for row in rows], ["Alice", "Bob", ""])

    def test_bad_configuration_fails_instead_of_silently_assigning(self):
        invalid = [
            "managers: []",
            "managers: Alice",
            "managers: [Alice, Alice]",
            "managers: ['']",
            "managers: [null]",
            "managers: [Alice]\nclubs: {Club: [Bob]}",
            "clubs: {Club: []}",
            "clubs: {Club: Alice}",
            "{}",
            "[Alice, Bob]",
        ]
        with tempfile.TemporaryDirectory() as directory:
            config = Path(directory) / "bad.yml"
            for source in invalid:
                config.write_text(source, encoding="utf-8")
                for root, builder in BUILDERS:
                    with self.subTest(root=root, source=source):
                        with self.assertRaises(ValueError):
                            builder.load_managers(config)

    def test_manager_column_is_written_and_validated_in_xlsx(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            for (root, builder), (_, validator) in zip(BUILDERS, VALIDATORS):
                with self.subTest(root=root):
                    config = root / "config/managers_by_club.yml"
                    rows = [{"client_id": f"{index:09}", "normalized_club": "Новый клуб"} for index in range(100)]
                    builder.assign_managers(rows, builder.load_managers(config))
                    template = directory / "template.xlsx"
                    workbook = Workbook()
                    workbook.active.append(builder.BASE_MAIN_HEADERS)
                    workbook.active.append(builder.BASE_MAIN_RUS_HEADERS)
                    workbook.save(template)
                    workbook.close()
                    output = directory / "clients.xlsx"
                    builder.write_main_xlsx(template, output, rows)
                    workbook = load_workbook(output, read_only=True, data_only=True)
                    actual = list(workbook.active.iter_rows(min_row=3, max_col=10, values_only=True))
                    workbook.close()
                    self.assertEqual({row[8] for row in actual}, EXPECTED_MANAGERS)
                    self.assertEqual(validator.validate_manager_assignments(actual, rows, config), [])
                    changed = [list(row) for row in actual]
                    changed[0][8] = next(name for name in EXPECTED_MANAGERS if name != changed[0][8])
                    self.assertTrue(validator.validate_manager_assignments(changed, rows, config))
                    changed[0][8] = "Старый менеджер"
                    self.assertTrue(validator.validate_manager_assignments(changed, rows, config))
                    changed[0][8] = ""
                    self.assertTrue(validator.validate_manager_assignments(changed, rows, config))

    def test_legacy_validator_accepts_global_configuration(self):
        path = ROOT.parent / "scripts/13_validate_part2_outputs.py"
        if not path.exists():
            self.skipTest("Historical validator is outside standalone package")
        validator = load_module("legacy_manager_validator", path)
        self.assertEqual(
            validator.load_managers(ROOT / "config/managers_by_club.yml"),
            {"*": EXPECTED_MANAGERS},
        )


if __name__ == "__main__":
    unittest.main()
