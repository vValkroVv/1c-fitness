#!/usr/bin/env python3
"""Independently reconcile the photo ZIP with the delivered client workbook."""

import argparse
from collections import Counter
import csv
import hashlib
import io
import json
from pathlib import Path
import re
import zipfile

from openpyxl import load_workbook


def phones(value):
    result = set()
    for part in re.split(r"[,;]", str(value or "")):
        digits = re.sub(r"\D", "", part)
        if len(digits) == 10:
            result.add("7" + digits)
        elif len(digits) == 11 and digits[0] in "78":
            result.add("7" + digits[1:])
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--delivery", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    root = args.delivery
    photo_report = json.loads((root / "reports/photos.json").read_text())
    delivery_manifest = json.loads((root / "reports/delivery_manifest.json").read_text())
    assert photo_report["status"] == "PASS"
    assert photo_report["cutoff_at"] == delivery_manifest["cutoff_contract"]["cutoff_at"]
    assert photo_report["backup_finish_at"] == delivery_manifest["cutoff_contract"]["backup_finish_at"]
    workbook_path = next(root.glob("fitbase_active_clients_import_zayavki_*.xlsx"))
    assert hashlib.sha256(workbook_path.read_bytes()).hexdigest() == photo_report["clients_xlsx_sha256"]
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = list(next(sheet.iter_rows(values_only=True)))
    clients = {}
    for values in sheet.iter_rows(min_row=3, values_only=True):
        row = dict(zip(headers, values))
        client_id = str(row["client_id"])
        assert client_id not in clients
        clients[client_id] = row
    workbook.close()
    assert len(clients) == photo_report["delivered_clients"]
    archive_path = next(root.glob("fitbase_client_photos_*.zip"))
    inner = photo_report["inner_directory"]
    with zipfile.ZipFile(archive_path) as archive:
        names = archive.namelist()
        assert len(names) == len(set(names))

        def rows(name):
            with archive.open(f"{inner}/_reports/{name}") as handle:
                return list(csv.DictReader(io.TextIOWrapper(handle, encoding="utf-8-sig")))

        manifest = rows("manifest.csv")
        excluded = rows("excluded_clients.csv")
        exceptions = rows("phone_filename_exceptions.csv")
        photos = {name for name in names if "/photos/" in name and not name.endswith("/")}
        assert photos == {f"{inner}/photos/{row['filename']}" for row in manifest}
        seen_clients = set()
        suffix_rows = set()
        for row in manifest:
            client_id = row["client_id"]
            assert client_id in clients and client_id not in seen_clients
            seen_clients.add(client_id)
            client = clients[client_id]
            assert client["funnel"] in {"Действующие абонементы", "Реактивация"}
            assert row["funnel"] == client["funnel"]
            phone = row["assigned_phone"]
            assert re.fullmatch(r"7\d{10}", phone) and phone in phones(client["phone"])
            assert row["exported_phone"] == str(client["phone"] or "")
            assert row["filename"] in {f"{phone}.jpg", f"{phone}__{client_id}.jpg"}
            if row["filename"] != f"{phone}.jpg":
                suffix_rows.add(row["filename"])
            assert row["output_extension"] == "jpg"
            data = archive.read(f"{inner}/photos/{row['filename']}")
            # Reading the entire ZIP member also checks its CRC. The production
            # validator decoded every JPEG; matching hashes retain that proof.
            assert data[:2] == b"\xff\xd8"
            assert len(data) == int(row["bytes"])
            assert hashlib.sha256(data).hexdigest() == row["sha256"]
        assert suffix_rows <= {row["filename"] for row in exceptions}
        excluded_delivered = [row for row in excluded if row["client_id"] in clients]
        excluded_ids = {row["client_id"] for row in excluded_delivered}
        assert len(excluded_ids) == len(excluded_delivered)
        assert not (seen_clients & excluded_ids)
        assert seen_clients | excluded_ids == set(clients)
        assert len(manifest) == photo_report["build"]["photos"]
        for row in excluded_delivered:
            reason = row["reason"]
            client = clients[row["client_id"]]
            assert reason in {"missing_phone", "invalid_phone", "missing_or_unsupported_photo"}
            if reason == "missing_phone":
                # A punctuation/text placeholder also has no phone number.
                assert not re.search(r"\d", str(client["phone"] or ""))
            if reason == "invalid_phone":
                assert not phones(client["phone"])
    result = {
        "status": "PASS", "cutoff_at": photo_report["cutoff_at"],
        "delivered_clients": len(clients), "jpeg_files": len(manifest),
        "photo_funnels": dict(Counter(row["funnel"] for row in manifest)),
        "delivered_clients_without_photo": len(excluded_ids),
        "delivered_exclusion_reasons": dict(Counter(row["reason"] for row in excluded_delivered)),
        "non_delivered_clients_excluded": len(excluded) - len(excluded_delivered),
        "phone_suffix_filenames": len(suffix_rows),
        "checks": "All JPEG member sets, IDs, funnels, phones, filenames, bytes, SHA-256, CRC and complete client partition",
    }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
