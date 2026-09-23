"""Keep template prices numeric and prevent the Fitbase thousands-prefix regression."""

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]


def load_builder(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


class TemplateNumberFormatTests(unittest.TestCase):
    def test_template_exports_preserve_amounts_without_grouping(self):
        for base, label in [(ROOT, "pipeline"), (ROOT.parent, "legacy")]:
            for kind, filename in [
                ("membership", "19_build_membership_import_xlsx.py"),
                ("service", "23_build_services_import_xlsx.py"),
            ]:
                with self.subTest(builder=label, kind=kind), tempfile.TemporaryDirectory() as directory:
                    builder = load_builder(base / "scripts" / filename, f"plain_{label}_{kind}")
                    template, output = Path(directory) / "template.xlsx", Path(directory) / "output.xlsx"
                    workbook = Workbook()
                    sheet = workbook.active
                    sheet.append(builder.TEMPLATE_HEADERS)
                    sheet.append(builder.TEMPLATE_RUS_HEADERS)
                    for col in range(1, len(builder.TEMPLATE_HEADERS) + 1):
                        sheet.cell(3, col).number_format = "#,##0.00"
                    workbook.save(template)
                    workbook.close()
                    amounts = [11990, 15000, 0, 12.5]
                    rows = [{"name": f"Template {i}", "price": amount, "duration": 12,
                             "visits": 1000} for i, amount in enumerate(amounts)]
                    if kind == "membership":
                        builder.write_workbook(template, output, builder.TEMPLATE_HEADERS, rows,
                                               builder.TEMPLATE_RUS_HEADERS)
                    else:
                        output.write_bytes(template.read_bytes())
                        builder.write_workbook(output, builder.TEMPLATE_HEADERS,
                                               builder.TEMPLATE_RUS_HEADERS, rows)
                    result = load_workbook(output)
                    result_sheet = result.active
                    price_col = builder.TEMPLATE_HEADERS.index("price") + 1
                    self.assertEqual(result_sheet.max_row, len(rows) + 2)
                    self.assertEqual([result_sheet.cell(i + 3, price_col).value for i in range(4)], amounts)
                    for row in result_sheet.iter_rows(min_row=3):
                        for cell in row:
                            if isinstance(cell.value, (int, float)):
                                self.assertEqual(cell.data_type, "n")
                                self.assertEqual(cell.number_format, "General")
                    result.close()


if __name__ == "__main__":
    unittest.main()
